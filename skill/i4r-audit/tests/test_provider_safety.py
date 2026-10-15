"""The participant-code rule shared by Phase-1 admission and the coordinator (2.6.5).

Every test builds its deposit on disk and runs the real functions; nothing here greps source.
The video-speed audit (20260926-audit-8fcff20d) was lost to the rule this pins: its SPSS column
names INT_P05 and DUR_P15 have the participant-code shape, the coordinator refused the five
Phase-1 reports that quoted them, and Phase 1 had already admitted and paid for those reports.
"""
from __future__ import annotations

import gzip
import hashlib
import io
import json
import os
import pathlib
import pickle
import random
import shutil
import struct
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import zipfile
import zlib
from unittest import mock

SKILL = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))

import phase1_runtime  # noqa: E402
import provider_safety  # noqa: E402
import run_coordinator  # noqa: E402
import build_draft_registry  # noqa: E402
import build_safe_evidence  # noqa: E402
import run_verifier  # noqa: E402


# ------------------------------------------------------------------------------------ fixtures
def _cells(value, width: int) -> list[bytes]:
    if width == 0:
        return [struct.pack("<d", float(value))]
    raw = str(value).encode("ascii").ljust(((width + 7) // 8) * 8, b" ")
    return [raw[index:index + 8] for index in range(0, len(raw), 8)]


def write_sav(path: pathlib.Path, variables: list[tuple[str, int]], rows: list[list],
              *, compressed: bool = False, long_names: dict[str, str] | None = None,
              value_labels: dict[str, list[tuple[object, str]]] | None = None,
              label_variable_record: bool = True) -> None:
    """A minimal SPSS system file built from the published record layout.

    ``variables`` holds (short name, width); width 0 is numeric, a string wider than eight bytes
    gets one continuation record per extra 8-byte segment, exactly as SPSS writes it.
    ``value_labels`` maps a short name to (value, label) pairs, written as a record 3 (count; per
    label an 8-byte value, a length byte and the label padded so 1 + length is a multiple of 8)
    and the record 4 that must follow it (count, then 1-based dictionary indices). R haven reads
    these files back with the same labels. ``label_variable_record=False`` omits the record 4.
    """
    nominal = sum(1 if width == 0 else (width + 7) // 8 for _, width in variables)
    bias = 100.0
    header = (
        b"$FL2" + b"@(#) SPSS DATA FILE i4r test".ljust(60)
        + struct.pack("<iiiii", 2, nominal, 1 if compressed else 0, 0, len(rows))
        + struct.pack("<d", bias) + b"01 Jan 26" + b"00:00:00" + b" " * 64 + b"\0\0\0"
    )
    assert len(header) == 176
    body = bytearray(header)
    index: dict[str, int] = {}
    position = 0
    for name, width in variables:
        position += 1
        index[name] = position
        body += struct.pack("<iiiiii", 2, width, 0, 0, 0x050800, 0x050800)
        body += name.upper().encode("ascii").ljust(8)[:8]
        for _ in range(((width + 7) // 8) - 1 if width else 0):
            position += 1
            body += struct.pack("<iiiiii", 2, -1, 0, 0, 0, 0) + b" " * 8
    widths = dict(variables)
    for name, pairs in (value_labels or {}).items():
        body += struct.pack("<ii", 3, len(pairs))
        for value, label in pairs:
            raw = label.encode("utf-8")
            body += (struct.pack("<d", float(value)) if widths[name] == 0
                     else str(value).encode("ascii").ljust(8)[:8])
            body += bytes([len(raw)]) + raw.ljust(((len(raw) + 1 + 7) // 8) * 8 - 1, b" ")
        if label_variable_record:
            body += struct.pack("<iii", 4, 1, index[name])
    if long_names:
        text = "\t".join(f"{short}={long}" for short, long in long_names.items()).encode("utf-8")
        body += struct.pack("<iiii", 7, 13, 1, len(text)) + text
    body += struct.pack("<ii", 999, 0)
    cells: list[tuple[bytes, bool]] = []  # (8 bytes, is_numeric)
    for row in rows:
        for (name, width), value in zip(variables, row):
            for cell in _cells(value, width):
                cells.append((cell, width == 0))
    if not compressed:
        body += b"".join(cell for cell, _ in cells)
    else:
        codes: list[int] = []
        raws: list[bytes] = []
        for cell, numeric in cells:
            if numeric and float(struct.unpack("<d", cell)[0]).is_integer() \
                    and -99 <= struct.unpack("<d", cell)[0] <= 151:
                codes.append(int(struct.unpack("<d", cell)[0] + bias))
                raws.append(b"")
            else:
                codes.append(253)
                raws.append(cell)
        codes.append(252)
        raws.append(b"")
        for start in range(0, len(codes), 8):
            block = codes[start:start + 8]
            body += bytes(block + [0] * (8 - len(block)))
            body += b"".join(raws[start:start + 8])
    path.write_bytes(bytes(body))


def write_dta118(path: pathlib.Path, names: list[str], text_values: list[str]) -> None:
    """The Stata 118 sections the parser reads (header, variable names, data), per the format spec."""
    varnames = b"".join(name.encode("ascii").ljust(129, b"\0") for name in names)
    data = b"".join(value.encode("ascii").ljust(16, b"\0") for value in text_values)
    path.write_bytes(
        b"<stata_dta><header><release>118</release><byteorder>LSF</byteorder>"
        + b"<K>" + struct.pack("<H", len(names)) + b"</K>"
        + b"<N>" + struct.pack("<Q", 1) + b"</N></header><map>" + b"\0" * 112 + b"</map>"
        + b"<varnames>" + varnames + b"</varnames><data>" + data + b"</data></stata_dta>"
    )


# The header saveRDS writes (XDR format 3, written by R 4.4.2, readable from 3.5.0, UTF-8).
R_XDR_HEADER = b"X\n" + struct.pack(">iiii", 3, 0x040402, 0x030500, 5) + b"UTF-8"


def r_strings(*values: str) -> bytes:
    """An R character vector in XDR serialization: a STRSXP holding one CHARSXP per value."""
    body = struct.pack(">ii", 16, len(values))
    for value in values:
        raw = value.encode("ascii")
        body += struct.pack(">ii", 0x00040009, len(raw)) + raw
    return body


def zip_bytes(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


class Deposit:
    def __init__(self) -> None:
        self._tmp = tempfile.TemporaryDirectory(prefix="i4r-provider-safety-")
        self.root = pathlib.Path(self._tmp.name) / "audit"
        self.pkg = self.root / "pkg"
        self.pkg.mkdir(parents=True)

    def close(self) -> None:
        self._tmp.cleanup()


class DepositTest(unittest.TestCase):
    def setUp(self) -> None:
        self.deposit = Deposit()
        self.pkg = self.deposit.pkg

    def tearDown(self) -> None:
        self.deposit.close()

    def vocabulary(self) -> provider_safety.DepositVocabulary:
        return provider_safety.deposit_vocabulary(self.pkg)


# ---------------------------------------------------------------------------- the vocabulary
class DepositVocabularyTests(DepositTest):
    def test_csv_header_names_are_exempt_and_identifier_values_are_not(self) -> None:
        (self.pkg / "data.csv").write_text("id,INT_P05,DUR_P15\nSUBJ_P12,3.2,1.1\nSUBJ_P13,2.9,1.4\n")
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05", "DUR_P15"}))
        self.assertIn("SUBJ_P12", vocabulary.values)

    def test_a_header_that_is_also_a_data_value_is_not_exempt(self) -> None:
        # Negative control: a column that happens to share a real code's spelling admits nothing.
        (self.pkg / "a.csv").write_text("id,SUBJ_P12\n1,4\n")
        (self.pkg / "b.csv").write_text("participant,score\nSUBJ_P12,4\n")
        self.assertEqual(self.vocabulary().exempt, frozenset())

    def test_one_column_per_participant_admits_none_of_its_codes(self) -> None:
        # A transposed export: participants are the columns, so its headers ARE participant codes.
        (self.pkg / "curves.csv").write_text(
            "time,SUBJ_P1,SUBJ_P2,SUBJ_P3,SUBJ_P4,SUBJ_P5\n0.0,1,2,3,4,5\n0.1,2,3,4,5,6\n")
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.participant_columns,
                         frozenset({"SUBJ_P1", "SUBJ_P2", "SUBJ_P3", "SUBJ_P4", "SUBJ_P5"}))
        self.assertEqual(vocabulary.exempt, frozenset())
        # A variable dictionary with a few code-shaped names is still admitted.
        (self.pkg / "study.csv").write_text(
            "id,age,sex,INT_P05,INT_P10,INT_P15,INT_P20,rating,group,condition\n"
            "1,20,f,1,2,3,4,5,a,b\n")
        self.assertEqual(self.vocabulary().exempt,
                         frozenset({"INT_P05", "INT_P10", "INT_P15", "INT_P20"}))

    def test_a_crashing_parser_makes_the_file_unreadable(self) -> None:
        # A value-label record truncated mid-dictionary.
        data = bytearray()
        write_sav(self.pkg / "ok.sav", [("INT_P05", 0)], [[1]])
        good = (self.pkg / "ok.sav").read_bytes()
        (self.pkg / "ok.sav").unlink()
        cut = good.index(struct.pack("<ii", 999, 0))
        data = good[:cut] + struct.pack("<ii", 3, 5) + b"\1"
        (self.pkg / "bad.sav").write_bytes(bytes(data))
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unreadable, ("bad.sav",))
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_semicolon_tables_are_read_and_prose_is_not(self) -> None:
        (self.pkg / "data.txt").write_text("INT_P05;INT_N05\n1;2\n3;4\n")
        (self.pkg / "README.txt").write_text(
            "Column INT_P20 holds ratings, and row SUBJ_P12 was excluded before sharing.\n"
            "See the syntax for details.\n"
        )
        self.assertEqual(self.vocabulary().names, frozenset({"INT_P05"}))

    def test_code_and_syntax_files_contribute_no_exemption(self) -> None:
        (self.pkg / "analysis.sps").write_text("COMPUTE INT_P05 = MEAN(a, b).\nSELECT IF id <> 'SUBJ_P12'.\n")
        self.assertEqual(self.vocabulary().names, frozenset())

    def test_spss_dictionary_short_and_long_names(self) -> None:
        write_sav(self.pkg / "study.sav", [("VP", 8), ("INT_P05", 0), ("DURLONGP", 0)],
                  [["A01", 3, 2], ["A02", 4, 1]],
                  long_names={"VP": "VP", "INT_P05": "INT_P05", "DURLONGP": "DUR_P15"})
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unreadable, ())
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05", "DUR_P15"}))

    def test_spss_bytecode_value_split_across_a_command_block_is_still_found(self) -> None:
        # Seven compressible numbers then a 16-byte string: the string's two cells straddle the
        # boundary between two command blocks, so its raw bytes are not contiguous on disk.
        variables = [(f"N{index}", 0) for index in range(7)] + [("CODE", 16), ("INT_P05", 0)]
        write_sav(self.pkg / "study.sav", variables,
                  [[1, 2, 3, 4, 5, 6, 7, "ABCDEF_P1234567", 3]], compressed=True)
        raw = (self.pkg / "study.sav").read_bytes()
        self.assertNotIn(b"ABCDEF_P1234567", raw, "fixture must split the value on disk")
        vocabulary = self.vocabulary()
        self.assertIn("ABCDEF_P1234567", vocabulary.values)
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))

    def test_spss_value_matching_a_dictionary_name_revokes_the_exemption(self) -> None:
        write_sav(self.pkg / "study.sav", [("CODE", 8), ("INT_P05", 0)], [["INT_P05", 3]])
        self.assertEqual(self.vocabulary().exempt, frozenset())

    def test_xlsx_first_row_is_the_dictionary(self) -> None:
        import openpyxl
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["id", "INT_P05"])
        sheet.append(["SUBJ_P12", 3])
        workbook.save(self.pkg / "data.xlsx")
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))
        self.assertIn("SUBJ_P12", vocabulary.values)

    def test_stata_118_variable_names(self) -> None:
        write_dta118(self.pkg / "data.dta", ["id", "INT_P05"], ["SUBJ_P12"])
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))
        self.assertIn("SUBJ_P12", vocabulary.values)

    def test_an_unreadable_data_file_contributes_no_names(self) -> None:
        (self.pkg / "broken.sav").write_bytes(b"$FL2" + b"\0" * 40)
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.names, frozenset())
        self.assertEqual(vocabulary.unreadable, ("broken.sav",))

    def test_missing_package_admits_nothing(self) -> None:
        self.assertEqual(provider_safety.deposit_exemptions(self.pkg / "absent"), frozenset())
        self.assertEqual(provider_safety.deposit_exemptions(None), frozenset())

    def test_cache_notices_a_changed_file(self) -> None:
        path = self.pkg / "data.csv"
        path.write_text("id,INT_P05\n1,2\n")
        self.assertEqual(self.vocabulary().exempt, frozenset({"INT_P05"}))
        path.write_text("id,INT_P05,DUR_P15\n1,2,3\n")
        self.assertEqual(self.vocabulary().exempt, frozenset({"INT_P05", "DUR_P15"}))


# ------------------------------------------------------------------------------- the matcher
class FirstParticipantCodeTests(unittest.TestCase):
    def test_raw_shape_is_refused_unless_exempt(self) -> None:
        text = "Line one.\nThe INT_P05 mean is 3.2.\n"
        self.assertEqual(provider_safety.first_participant_code(text), provider_safety.CodeHit(2, "raw"))
        self.assertIsNone(provider_safety.first_participant_code(text, {"INT_P05"}))

    def test_an_exempt_token_does_not_hide_a_later_real_code(self) -> None:
        text = "INT_P05 is a column.\nSUBJ_P12 is a participant.\n"
        self.assertEqual(provider_safety.first_participant_code(text, {"INT_P05"}),
                         provider_safety.CodeHit(2, "raw"))

    def test_row_signature_is_refused_unless_its_code_is_exempt(self) -> None:
        text = "header\nP12, 3.4, 5\n"
        self.assertEqual(provider_safety.first_participant_code(text), provider_safety.CodeHit(2, "row"))
        self.assertIsNone(provider_safety.first_participant_code(text, {"P12"}))

    def test_the_patterns_are_the_coordinators_historical_ones(self) -> None:
        self.assertIs(run_coordinator.RAW_PARTICIPANT_CODE, provider_safety.RAW_PARTICIPANT_CODE)
        self.assertIs(run_coordinator.ROW_PARTICIPANT_CODE, provider_safety.ROW_PARTICIPANT_CODE)
        for text in ("SA_P1 was seen", "Unsafe row: SA_P1,2022-05-09,118,74", "\nP7,1.5\n"):
            self.assertTrue(provider_safety.has_participant_code(text), text)
        for text in ("TLI equals 0.9546", "Table P1 of the appendix", "item_p05 lower case"):
            self.assertFalse(provider_safety.has_participant_code(text), text)


# ------------------------------------------------------------------- both stages, same verdict
class BothStagesAgreeTests(DepositTest):
    REPORT = "## Findings\nThe INT_P05 and DUR_P15 means reproduce.\n"

    def setUp(self) -> None:
        super().setUp()
        (self.pkg / "study.csv").write_text("id,INT_P05,DUR_P15\nSUBJ_P12,3.2,1.1\n")

    def phase_one(self, report: str, result: dict | None = None, *, exempt=None) -> str | None:
        exemptions = self.vocabulary().exempt if exempt is None else exempt
        try:
            phase1_runtime.check_output_hygiene(report, result or {"scientific_fields": {}}, exemptions)
        except phase1_runtime.OutputValidationError as exc:
            return exc.rule
        return None

    def coordinator(self, text: str) -> bool:
        try:
            run_coordinator.validate_provider_visible_string(
                text, "Phase-1 report", run_coordinator.deposit_exemptions(self.deposit.root))
        except run_coordinator.CoordinatorError:
            return False
        return True

    def test_column_names_pass_both_stages(self) -> None:
        self.assertIsNone(self.phase_one(self.REPORT))
        self.assertTrue(self.coordinator(self.REPORT))

    def test_without_the_deposit_both_stages_refuse(self) -> None:
        # The pre-2.6.5 coordinator verdict, now also Phase 1's: no exemption without evidence.
        self.assertEqual(self.phase_one(self.REPORT, exempt=frozenset()), "report_participant_code")
        with self.assertRaises(run_coordinator.CoordinatorError) as caught:
            run_coordinator.validate_provider_visible_string(self.REPORT, "Phase-1 report")
        self.assertIn("not provider-safe", str(caught.exception))

    def test_a_real_code_is_refused_by_both_stages(self) -> None:
        report = self.REPORT + "The row for SUBJ_P12 is duplicated.\n"
        self.assertEqual(self.phase_one(report), "report_participant_code")
        self.assertFalse(self.coordinator(report))

    def test_a_code_in_result_json_is_refused_at_phase_one(self) -> None:
        result = {"scientific_fields": {"findings": [{"claim": "SUBJ_P12 is duplicated"}]}}
        self.assertEqual(self.phase_one(self.REPORT, result), "result_participant_code")

    def test_phase_one_rejection_is_not_offered_a_focused_repair(self) -> None:
        with self.assertRaises(phase1_runtime.OutputValidationError) as caught:
            phase1_runtime.check_output_hygiene("SUBJ_P12\n", {}, frozenset())
        self.assertFalse(caught.exception.repairable)

    def test_snapshot_exemptions_read_raw_pkg(self) -> None:
        snapshot = self.deposit.root / "snapshot"
        (snapshot / "raw").mkdir(parents=True)
        (snapshot / "raw" / "pkg").symlink_to(self.pkg, target_is_directory=True)
        self.assertEqual(phase1_runtime.snapshot_deposit_exemptions(snapshot),
                         frozenset({"INT_P05", "DUR_P15"}))

    def test_code_beside_a_date_keeps_the_2_6_4_rule_in_both_output_checks(self) -> None:
        # The exemption covers the bare-token scan only. A code-shaped token beside a date is still
        # refused by the coordinator's output check AND by check_consistency.sh on the same file,
        # so the post-call check and the post-delta gate cannot disagree.
        line = "INT_P05 was recoded on 2024-01-02\n"
        with self.assertRaises(run_coordinator.CoordinatorError):
            run_coordinator.validate_output_privacy(line, "x")
        report = self.deposit.root / "report.md"
        report.write_text("## Abstract\n\n" + line)
        result = subprocess.run(["bash", str(SKILL / "scripts/check_consistency.sh"), str(report)],
                                capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("participant code and date on one line", result.stdout)


class ParticipantColumnTests(DepositTest):
    """One column per participant must never be exempt, however the layout is dressed up."""

    def test_three_participant_columns_are_enough(self) -> None:
        (self.pkg / "scd.csv").write_text("session,ATH_P1,ATH_P2,ATH_P3\n1,4,5,6\n2,5,6,7\n")
        self.assertEqual(self.vocabulary().exempt, frozenset())

    def test_two_code_shaped_variables_are_still_a_dictionary(self) -> None:
        (self.pkg / "study.csv").write_text("id,INT_P05,DUR_P15\n1,3.2,1.1\n")
        self.assertEqual(self.vocabulary().exempt, frozenset({"INT_P05", "DUR_P15"}))

    def test_blank_header_cells_do_not_dilute_the_share(self) -> None:
        # A merged L/R header exported to CSV leaves one blank cell per participant.
        (self.pkg / "lr.csv").write_text(
            "time,SUBJ_P01,,SUBJ_P02,,SUBJ_P03,,SUBJ_P04,\n0,1,2,3,4,5,6,7,8\n1,2,3,4,5,6,7,8,9\n")
        self.assertEqual(self.vocabulary().exempt, frozenset())

    def test_companion_columns_count_as_participant_columns(self) -> None:
        (self.pkg / "pairs.csv").write_text(
            "time,SUBJ_P01,SUBJ_P01_SD,SUBJ_P02,SUBJ_P02_SD,SUBJ_P03,SUBJ_P03_SD\n0,1,2,3,4,5,6\n")
        vocabulary = self.vocabulary()
        self.assertTrue({"SUBJ_P01", "SUBJ_P02", "SUBJ_P03"} <= vocabulary.participant_columns)
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_each_workbook_sheet_is_judged_on_its_own(self) -> None:
        import openpyxl
        workbook = openpyxl.Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["id"] + [f"var{index}" for index in range(59)] + ["INT_P05"])
        summary.append([1] + [0] * 60)
        curves = workbook.create_sheet("Curves")
        curves.append(["time"] + [f"SUBJ_P{index:02d}" for index in range(1, 21)])
        curves.append([0.0] + [1.0] * 20)
        workbook.save(self.pkg / "data.xlsx")
        vocabulary = self.vocabulary()
        self.assertIn("SUBJ_P07", vocabulary.participant_columns)
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))


class ValueBackstopTests(DepositTest):
    """A token that is a data value anywhere in the deposit is never exempt."""

    def setUp(self) -> None:
        super().setUp()
        (self.pkg / "study.csv").write_text("id,SUBJ_P12,INT_P05\n1,2,3\n")

    def exempt(self) -> frozenset[str]:
        return self.vocabulary().exempt

    def test_ids_in_a_gzipped_rds_revoke_the_exemption(self) -> None:
        (self.pkg / "ids.rds").write_bytes(gzip.compress(R_XDR_HEADER + r_strings("SUBJ_P12", "SUBJ_P13")))
        self.assertEqual(self.exempt(), frozenset({"INT_P05"}))

    def test_an_r_string_whose_length_byte_is_a_letter_is_still_found(self) -> None:
        # A 65-byte string is preceded by the length byte 0x41 ("A"), which glues onto its first
        # token; the raw substring search still finds the candidate name.
        note = "SUBJ_P12 withdrew after session two; asked for the data deletion."
        self.assertEqual(len(note), 65)
        (self.pkg / "notes.rds").write_bytes(gzip.compress(R_XDR_HEADER + r_strings(note)))
        self.assertEqual(self.exempt(), frozenset({"INT_P05"}))

    def test_ids_in_a_zipped_csv_revoke_the_exemption(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("raw/ids.csv", "participant\nSUBJ_P12\n")
            archive.writestr("videos/clip.mp4", b"SUBJ_P99 in media bytes is ignored")
        (self.pkg / "archive.zip").write_bytes(buffer.getvalue())
        self.assertEqual(self.exempt(), frozenset({"INT_P05"}))

    def test_ids_in_a_one_column_list_revoke_the_exemption(self) -> None:
        (self.pkg / "ids.txt").write_text("SUBJ_P12\nSUBJ_P13\n")
        self.assertEqual(self.exempt(), frozenset({"INT_P05"}))

    def test_an_unreadable_data_file_disables_every_exemption(self) -> None:
        # Stata 114 is not read; its bytes are no fallback (2.6.5 review), so nothing is exempt.
        (self.pkg / "old.dta").write_bytes(b"\x72\x02\x01\x00" + b"SUBJ_P12".ljust(33, b"\0"))
        vocabulary = self.vocabulary()
        self.assertTrue(vocabulary.unscannable)
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_utf16_text_is_decoded(self) -> None:
        (self.pkg / "export.txt").write_bytes("participant\tscore\r\nSUBJ_P12\t3\r\n".encode("utf-16"))
        self.assertEqual(self.exempt(), frozenset({"INT_P05"}))

    def test_a_file_that_cannot_be_scanned_disables_every_exemption(self) -> None:
        (self.pkg / "blob.gz").write_bytes(b"\x1f\x8bnot really gzip")
        vocabulary = self.vocabulary()
        self.assertTrue(vocabulary.unscannable)
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_code_documents_and_media_are_not_values(self) -> None:
        (self.pkg / "analysis.R").write_text("d <- subset(d, id != 'SUBJ_P12'); mean(d$INT_P05)\n")
        (self.pkg / "README.md").write_text("INT_P05 is the 0.5 s intention rating.\n")
        (self.pkg / "clip.mp4").write_bytes(b"INT_P05")
        self.assertEqual(self.exempt(), frozenset({"SUBJ_P12", "INT_P05"}))

    def test_a_code_flush_against_the_next_binary_field_is_found(self) -> None:
        # Stata stores fixed-width strings; SUBJ_P12 filling its field abuts the next byte.
        write_dta118(self.pkg / "data.dta", ["id", "score"], [])
        path = self.pkg / "data.dta"
        path.write_bytes(path.read_bytes().replace(b"<data></data>", b"<data>SUBJ_P127ABC</data>"))
        self.assertEqual(self.exempt(), frozenset({"INT_P05"}))


class RobustnessTests(DepositTest):
    def test_negative_spss_record_length_does_not_hang(self) -> None:
        write_sav(self.pkg / "ok.sav", [("INT_P05", 0)], [[1]])
        data = bytearray((self.pkg / "ok.sav").read_bytes())
        cut = data.index(struct.pack("<ii", 999, 0))
        bad = bytes(data[:cut]) + struct.pack("<ii", 4, -2) + bytes(data[cut:])
        (self.pkg / "ok.sav").write_bytes(bad)
        code = (f"import sys; sys.path.insert(0, {str(SKILL / 'scripts')!r}); import provider_safety as p; "
                f"v = p.deposit_vocabulary({str(self.pkg)!r}); print(v.unreadable)")
        started = time.monotonic()
        result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=60)
        self.assertLess(time.monotonic() - started, 30)
        self.assertIn("ok.sav", result.stdout)

    def test_bytecode_numbers_cost_one_byte(self) -> None:
        stream = bytes([101] * 8) * 1000 + bytes([252] + [0] * 7)
        self.assertEqual(len(provider_safety._spss_bytecode(stream)), 8000)

    def test_value_scan_is_linear_on_long_underscore_runs(self) -> None:
        payload = ("AB_" * 60000 + "P1").encode()
        started = time.monotonic()
        tokens = provider_safety._value_tokens(payload)
        self.assertLess(time.monotonic() - started, 10)
        self.assertIn("AB_P1", tokens)

    def test_concurrent_callers_compute_once(self) -> None:
        (self.pkg / "study.csv").write_text("id,INT_P05\n1,2\n")
        calls = []
        original = provider_safety._compute
        def counting(*args, **kwargs):
            calls.append(1)
            time.sleep(0.2)
            return original(*args, **kwargs)
        provider_safety._vocabulary.cache_clear()
        provider_safety._compute = counting
        try:
            threads = [threading.Thread(target=provider_safety.deposit_vocabulary, args=(self.pkg,))
                       for _ in range(5)]
            for thread in threads: thread.start()
            for thread in threads: thread.join()
        finally:
            provider_safety._compute = original
        self.assertEqual(len(calls), 1)

    def test_internal_symlinks_are_read_as_phase_one_copies_them(self) -> None:
        # Phase 1's snapshot copies an internal symlink's target under the link's name
        # (phase1_runtime.source_files); the coordinator reads the live pkg/. Both must agree.
        import shutil
        for target in ("export/real.csv", "export/study_table"):
            with self.subTest(target=target):
                pkg = self.deposit.root / f"pkg_{target.replace('/', '_')}"
                (pkg / "export").mkdir(parents=True)
                (pkg / target).write_text("id,INT_P05\n1,2\n")
                os.symlink(target, pkg / "study.csv")
                copied = self.deposit.root / f"snap_{target.replace('/', '_')}"
                shutil.copytree(pkg, copied, symlinks=False)
                live = provider_safety.deposit_exemptions(pkg)
                self.assertEqual(live, provider_safety.deposit_exemptions(copied))
                self.assertEqual(live, frozenset({"INT_P05"}) if target.endswith(".csv") else frozenset())


class DownstreamAgreementTests(DepositTest):
    def test_safe_evidence_receipt_uses_the_same_exemption(self) -> None:
        safe = self.deposit.root / "safe"
        safe.mkdir()
        (safe / "summary.md").write_text("The INT_P05 mean is 3.2; SUBJ_P12 was excluded.\n")
        patterns = build_safe_evidence.detectors()
        without = build_safe_evidence.scan_directory(safe, patterns)["participant_code"]
        with_exempt = build_safe_evidence.scan_directory(safe, patterns, frozenset({"INT_P05"}))["participant_code"]
        self.assertEqual((without, with_exempt), (2, 1))

    def test_draft_registry_never_cuts_a_token_into_a_code(self) -> None:
        text = "The HR_P95th value is high"
        self.assertEqual(build_draft_registry.clip(text, 9), "The")
        self.assertFalse(provider_safety.has_participant_code(build_draft_registry.clip(text, 9)))
        self.assertEqual(build_draft_registry.clip("short", 10), "short")

    def test_verifier_refuses_a_claim_it_would_have_to_repeat(self) -> None:
        row = {"id": "F9", "ledger_disposition": {"status": "INCLUDED"},
               "claim": "Participant SUBJ_P12 is duplicated in the data.",
               "provenance": {"origin_role": "critic_stats", "independently_found_by": ["critic_stats"],
                              "agreed_after_seeing": []}}
        with self.assertRaises(run_verifier.VerifierRefusal) as caught:
            run_verifier.select_hypotheses([row], ["F9"], ["auditor_reproduction"], {"source_files": []})
        self.assertIn("reword the claim", str(caught.exception.remedy) if hasattr(caught.exception, "remedy") else str(caught.exception))


class SpssValueLabelTests(DepositTest):
    """Value-label texts are data values (2.6.5 review P3/V4).

    A common SPSS layout keys participants by a number whose value labels are the codes. When the
    questionnaire items are named like codes too (P1..P20), the codes live only in the labels.
    """

    VARIABLES = [("ID", 0), ("AGE", 0), ("SEX", 0), ("P12", 0), ("RATING", 0), ("GROUP", 0), ("COND", 0)]
    ROWS = [[k, 20 + k % 5, 1 + k % 2, 3, 4, 1 + k % 2, 2] for k in range(1, 21)]
    CODES = [(k, f"P{k}") for k in range(1, 21)]

    def test_a_label_that_matches_a_dictionary_name_revokes_it(self) -> None:
        write_sav(self.pkg / "study.sav", self.VARIABLES, self.ROWS, value_labels={"ID": self.CODES})
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unscannable, ())
        self.assertIn("P12", vocabulary.names)
        self.assertIn("P12", vocabulary.values)
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_labels_that_are_not_codes_leave_the_name_exempt(self) -> None:
        # Negative control: the same file with ordinary labels admits its code-shaped name.
        write_sav(self.pkg / "study.sav", self.VARIABLES, self.ROWS,
                  value_labels={"GROUP": [(1, "control"), (2, "training")], "SEX": [(1, "f"), (2, "m")]})
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unscannable, ())
        self.assertEqual(vocabulary.exempt, frozenset({"P12"}))

    def test_labels_are_read_in_a_bytecode_compressed_file(self) -> None:
        write_sav(self.pkg / "study.sav", self.VARIABLES, self.ROWS, compressed=True,
                  value_labels={"ID": [(12, "participant P12 (withdrawn)")]})
        self.assertEqual(self.vocabulary().exempt, frozenset())

    def test_a_labelled_string_value_is_a_value(self) -> None:
        # For a string variable the 8-byte labelled value is stored text, present in the data or not.
        (self.pkg / "items.csv").write_text("id,SUBJ_P12,x,y\n1,2,3,4\n")
        write_sav(self.pkg / "codes.sav", [("CODE", 8), ("SCORE", 0)], [["A01", 3]],
                  value_labels={"CODE": [("SUBJ_P12", "control")]})
        vocabulary = self.vocabulary()
        self.assertIn("SUBJ_P12", vocabulary.values)
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_malformed_label_records_make_the_file_unscannable(self) -> None:
        write_sav(self.pkg / "good.sav", self.VARIABLES, self.ROWS, value_labels={"ID": self.CODES[:2]})
        good = (self.pkg / "good.sav").read_bytes()
        (self.pkg / "good.sav").unlink()
        record4 = struct.pack("<iii", 4, 1, 1)
        cut = good.index(struct.pack("<ii", 999, 0))
        cases = {
            "labels without their variable record": good.replace(record4, b""),
            "a variable record without labels": good[:cut] + record4 + good[cut:],
            "labels naming a variable not in the dictionary": good.replace(record4, struct.pack("<iii", 4, 1, 99)),
            "a truncated label": good[:good.index(b"P2") - 1],
        }
        for label, data in cases.items():
            with self.subTest(case=label):
                pkg = self.deposit.root / label.replace(" ", "_")
                pkg.mkdir()
                (pkg / "items.csv").write_text("id,INT_P05\n1,2\n")
                (pkg / "study.sav").write_bytes(data)
                vocabulary = provider_safety.deposit_vocabulary(pkg)
                self.assertIn("study.sav", vocabulary.unreadable)
                self.assertTrue(vocabulary.unscannable)
                self.assertEqual(vocabulary.exempt, frozenset())


class RWrittenFixtureTests(DepositTest):
    """Fixtures written by R itself: haven::write_sav, saveRDS and save. Skipped without R."""

    def rscript(self, code: str, *args: str) -> str:
        if shutil.which("Rscript") is None:
            self.skipTest("Rscript is not installed")
        result = subprocess.run(["Rscript", "-e", code, *args], capture_output=True, text=True, timeout=120)
        if "there is no package called" in result.stderr:
            self.skipTest("R package haven is not installed")
        self.assertEqual(result.returncode, 0, result.stderr)
        return result.stdout

    HAVEN = """
suppressMessages(library(haven))
a <- commandArgs(trailingOnly = TRUE)
n <- 20
d <- data.frame(age = 20 + seq_len(n) %% 5, P12 = rep(3, n), rating = rep(4, n),
                cond = rep(2, n), sex = rep(1, n))
if (a[2] == "codes") {
  d$id <- labelled(as.numeric(seq_len(n)), setNames(as.numeric(seq_len(n)), paste0("P", seq_len(n))))
} else if (a[2] == "long_string") {
  d$id <- labelled(sprintf("respondent%02d", seq_len(n)), c(P12 = "respondent12"))
} else {
  d$id <- labelled(as.numeric(seq_len(n)), c(control = 1, training = 2))
}
write_sav(d, a[1], compress = a[3])
back <- read_sav(a[1])
cat(paste(names(attr(back$id, "labels")), collapse = ","))
"""

    def test_haven_value_labels(self) -> None:
        expected = {"codes": ",".join(f"P{k}" for k in range(1, 21)),
                    "long_string": "P12", "words": "control,training"}
        for kind, compress in (("codes", "byte"), ("codes", "none"), ("codes", "zsav"),
                               ("long_string", "byte"), ("words", "byte")):
            with self.subTest(kind=kind, compress=compress):
                pkg = self.deposit.root / f"{kind}_{compress}"
                pkg.mkdir()
                name = "study.zsav" if compress == "zsav" else "study.sav"
                labels = self.rscript(self.HAVEN, str(pkg / name), kind, compress)
                self.assertEqual(labels, expected[kind], "haven must read the labels back")
                vocabulary = provider_safety.deposit_vocabulary(pkg)
                self.assertEqual(vocabulary.unscannable, ())
                self.assertIn("P12", vocabulary.names)
                self.assertEqual(vocabulary.exempt, frozenset() if kind != "words" else frozenset({"P12"}))

    def test_haven_reads_the_python_fixture_labels_back(self) -> None:
        write_sav(self.pkg / "study.sav", SpssValueLabelTests.VARIABLES, SpssValueLabelTests.ROWS,
                  compressed=True, value_labels={"ID": SpssValueLabelTests.CODES})
        labels = self.rscript("suppressMessages(library(haven)); d <- read_sav(commandArgs(TRUE)[1]); "
                              "cat(paste(names(attr(d$ID, 'labels')), collapse = ','))",
                              str(self.pkg / "study.sav"))
        self.assertEqual(labels, ",".join(f"P{k}" for k in range(1, 21)))

    def test_r_serialized_ids_are_values(self) -> None:
        (self.pkg / "items.csv").write_text("id,SUBJ_P12,INT_P05\n1,2,3\n")
        self.rscript("a <- commandArgs(TRUE); d <- data.frame(id = sprintf('SUBJ_P%02d', 1:40), rt = 1:40); "
                     "saveRDS(d, a[1]); save(d, file = a[2]); saveRDS(d, a[3], compress = 'xz')",
                     str(self.pkg / "ids.rds"), str(self.pkg / "ids.RData"), str(self.pkg / "ids_xz.rds"))
        for keep in ("ids.rds", "ids.RData", "ids_xz.rds"):
            with self.subTest(file=keep):
                pkg = self.deposit.root / keep.replace(".", "_")
                pkg.mkdir()
                shutil.copy(self.pkg / "items.csv", pkg)
                shutil.copy(self.pkg / keep, pkg)
                vocabulary = provider_safety.deposit_vocabulary(pkg)
                self.assertEqual(vocabulary.unscannable, ())
                self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))


class OpaqueFormatTests(DepositTest):
    """A value-bearing file whose bytes hide their text is unscannable (2.6.5 review)."""

    IDS = "id,rt\n" + "".join(f"SUBJ_P{k:02d},{500 + k}\n" for k in range(1, 41))

    def setUp(self) -> None:
        super().setUp()
        (self.pkg / "study.csv").write_text("id,SUBJ_P12,INT_P05\n1,2,3\n")

    def signed(self) -> dict[str, bytes]:
        noise = random.Random(7).randbytes(2048)
        return {
            "parquet": b"PAR1" + noise + b"PAR1",
            "arrow": b"ARROW1\0\0" + noise + b"ARROW1",
            "feather v1": b"FEA1" + noise + b"FEA1",
            "zstd": b"\x28\xb5\x2f\xfd" + noise,
            "lz4 frame": b"\x04\x22\x4d\x18" + noise,
            "snappy framing": b"\xff\x06\x00\x00sNaPpY" + noise,
            "7z": b"7z\xbc\xaf\x27\x1c" + noise,
            "rar": b"Rar!\x1a\x07\x00" + noise,
            "hdf5": b"\x89HDF\r\n\x1a\n" + noise,
            "hdf5 behind a user block": b"MATLAB 7.3 MAT-file".ljust(512, b" ") + b"\x89HDF\r\n\x1a\n" + noise,
            "raw zlib": zlib.compress(self.IDS.encode()),
            "matlab v5": b"MATLAB 5.0 MAT-file, Platform: MACI64".ljust(128, b" ") + noise,
            "pickle": pickle.dumps([f"SUBJ_P{k:02d}" for k in range(1, 41)], protocol=4),
            "sqlite": b"SQLite format 3\x00" + noise,
            "ole2 (.xls)": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + noise,
            "numpy": b"\x93NUMPY\x01\x00" + noise,
            "undecodable binary": noise,
        }

    def assert_unscannable(self, name: str, data: bytes) -> None:
        pkg = self.deposit.root / f"case_{len(list(self.deposit.root.iterdir()))}"
        pkg.mkdir()
        shutil.copy(self.pkg / "study.csv", pkg)
        (pkg / name).write_bytes(data)
        vocabulary = provider_safety.deposit_vocabulary(pkg)
        self.assertTrue(vocabulary.unscannable, name)
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_signed_opaque_formats_are_unscannable_under_any_name(self) -> None:
        for kind, data in self.signed().items():
            with self.subTest(kind=kind):
                self.assert_unscannable("ids.bin", data)

    def test_opaque_suffixes_and_binary_text_suffixes_are_unscannable(self) -> None:
        noise = random.Random(8).randbytes(4096)
        for name in ("ids.parquet", "ids.feather", "ids.mat", "ids.pkl", "ids.por", "ids.xls", "ids.dat",
                     "ids.txt", "ids.csv.gz"):
            with self.subTest(name=name):
                self.assert_unscannable(name, gzip.compress(noise) if name.endswith(".gz") else noise)

    def test_an_r_file_in_another_format_is_unscannable(self) -> None:
        self.assert_unscannable("ids.rds", b"\x28\xb5\x2f\xfd" + R_XDR_HEADER + r_strings("SUBJ_P12"))
        self.assert_unscannable("ids.RData", b"QS" + r_strings("SUBJ_P12"))

    def test_an_opaque_file_inside_an_archive_is_unscannable(self) -> None:
        self.assert_unscannable("data.zip", zip_bytes({"raw/ids.parquet": self.signed()["parquet"]}))
        self.assert_unscannable("data.zip", zip_bytes({"raw/ids.bin": self.signed()["raw zlib"]}))

    def test_text_is_still_read(self) -> None:
        # Negative controls: text of every kind stays scannable, and its values still count.
        (self.pkg / "ids.json").write_text(json.dumps([{"id": "SUBJ_P12"}]))
        (self.pkg / "params.csv").write_text("PAR1,PAR2,H\n1,2,3\n")  # opens like Parquet, is a table
        (self.pkg / "H.txt").write_text("H,W\n180,75\n")               # opens with a valid zlib header
        (self.pkg / "notes.txt").write_text("Überprüfung der Daten: 30 Teilnehmende, keine Ausfälle.\n")
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unscannable, ())
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))

    def test_folder_metadata_in_an_archive_is_read_as_before(self) -> None:
        # The video-speed deposit zips its stimuli with Finder metadata beside the clips.
        noise = random.Random(9).randbytes(512)
        archive = zip_bytes({
            "clips/.DS_Store": b"\x00\x00\x00\x01Bud1" + noise,
            "__MACOSX/._clips": b"\x00\x05\x16\x07\x00\x02\x00\x00" + noise,
            "__MACOSX/clips/._clip01.mp4": b"\x00\x05\x16\x07\x00\x02\x00\x00" + noise,
            "clips/clip01.mp4": noise,
        })
        (self.pkg / "stimuli.zip").write_bytes(archive)
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unscannable, ())
        self.assertEqual(vocabulary.exempt, frozenset({"SUBJ_P12", "INT_P05"}))

    def test_a_data_file_inside_an_archive_is_parsed(self) -> None:
        # The zipped .sav stores its code split across two bytecode command blocks; a byte scan
        # of the member would miss it.
        variables = [(f"N{index}", 0) for index in range(7)] + [("CODE", 16)]
        write_sav(self.pkg / "member.sav", variables, [[1, 2, 3, 4, 5, 6, 7, "SUBJ_P12"]], compressed=True)
        member = (self.pkg / "member.sav").read_bytes()
        (self.pkg / "member.sav").unlink()
        (self.pkg / "raw.zip").write_bytes(zip_bytes({"raw/ids.sav": member}))
        self.assertEqual(self.vocabulary().exempt, frozenset({"INT_P05"}))
        self.assert_unscannable("broken.zip", zip_bytes({"raw/ids.sav": member[:200]}))

    def test_the_backstop_runs_only_when_a_name_could_be_exempt(self) -> None:
        pkg = self.deposit.root / "no_candidates"
        pkg.mkdir()
        (pkg / "study.csv").write_text("id,age\n1,20\n")
        (pkg / "ids.parquet").write_bytes(self.signed()["parquet"])
        self.assertEqual(provider_safety.deposit_vocabulary(pkg).unscannable, ())


class StructuredReadFailureTests(DepositTest):
    """A structured reader's failure is a failed scan; raw bytes are no fallback (2.6.5 review)."""

    @staticmethod
    def zsav(trailer_offset: int | None) -> bytes:
        header = (b"$FL3" + b"@(#) SPSS DATA FILE i4r test".ljust(60)
                  + struct.pack("<iiiii", 2, 1, 2, 0, 1) + struct.pack("<d", 100.0)
                  + b"01 Jan 26" + b"00:00:00" + b" " * 64 + b"\0\0\0")
        dictionary = (header + struct.pack("<iiiiii", 2, 8, 0, 0, 0x010800, 0x010800) + b"ID      "
                      + struct.pack("<ii", 999, 0))
        payload = zlib.compress(bytes([253, 252, 0, 0, 0, 0, 0, 0]) + b"SUBJ_P12")
        offset = len(dictionary) + 24 + len(payload) if trailer_offset is None else trailer_offset
        return dictionary + struct.pack("<qqq", len(dictionary), offset, 48) + payload + b"\0" * 48

    def test_a_correct_zsav_is_read(self) -> None:
        (self.pkg / "items.csv").write_text("id,SUBJ_P12,x,y\n1,2,3,4\n")
        (self.pkg / "ids.zsav").write_bytes(self.zsav(None))
        vocabulary = self.vocabulary()
        self.assertIn("SUBJ_P12", vocabulary.values)
        self.assertEqual((vocabulary.unscannable, vocabulary.exempt), ((), frozenset()))

    def test_a_file_its_reader_rejects_disables_every_exemption(self) -> None:
        cases = {
            "ids.zsav": self.zsav(0),  # the trailer offset is wrong, so the data is never inflated
            "ids.sav": b"$FL2" + b"\0" * 300,
            "ids.dta": b"<stata_dta><header><release>118</release></header>",
            "ids.xlsx": b"PK\x03\x04garbage",
        }
        for name, data in cases.items():
            with self.subTest(file=name):
                pkg = self.deposit.root / name.replace(".", "_")
                pkg.mkdir()
                (pkg / "items.csv").write_text("id,INT_P05\n1,2\n")
                (pkg / name).write_bytes(data)
                vocabulary = provider_safety.deposit_vocabulary(pkg)
                self.assertIn(name, vocabulary.unreadable)
                self.assertTrue(any(entry.startswith(name) for entry in vocabulary.unscannable))
                self.assertEqual(vocabulary.exempt, frozenset())

    def test_a_workbook_without_openpyxl_is_unscannable(self) -> None:
        import openpyxl
        workbook = openpyxl.Workbook()
        workbook.active.append(["id", "INT_P05"])
        workbook.active.append(["SUBJ_P12", 3])
        workbook.save(self.pkg / "data.xlsx")
        (self.pkg / "items.csv").write_text("id,INT_P05\n1,2\n")
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            vocabulary = provider_safety.deposit_vocabulary(self.pkg)
        self.assertTrue(vocabulary.unscannable)
        self.assertEqual(vocabulary.exempt, frozenset())


class DepositWalkTests(DepositTest):
    """Every listing error fails closed; directory symlinks are still not followed (2.6.5 review)."""

    def setUp(self) -> None:
        super().setUp()
        self.locked: list[pathlib.Path] = []

    def tearDown(self) -> None:
        for path in self.locked:
            os.chmod(path, 0o755)
        super().tearDown()

    def lock(self, path: pathlib.Path) -> None:
        if hasattr(os, "geteuid") and os.geteuid() == 0:
            self.skipTest("root can list any directory")
        os.chmod(path, 0)
        self.locked.append(path)

    def test_a_directory_that_cannot_be_listed_disables_every_exemption(self) -> None:
        (self.pkg / "study.csv").write_text("x,SUBJ_P12\n1,2\n")
        secret = self.pkg / "secret"
        secret.mkdir()
        (secret / "ids.csv").write_text("id,score\nSUBJ_P12,1\n")
        self.lock(secret)
        vocabulary = self.vocabulary()
        self.assertTrue(any("cannot be listed" in entry for entry in vocabulary.unscannable))
        self.assertEqual(vocabulary.exempt, frozenset())

    def test_phase_one_refuses_a_directory_it_cannot_list(self) -> None:
        # Phase 1's snapshot would otherwise drop the directory silently and admit what the
        # coordinator, reading the live pkg/, refuses.
        (self.deposit.root / "paper").mkdir()
        (self.deposit.root / "paper" / "paper.txt").write_text("A paper.\n")
        (self.pkg / "study.csv").write_text("x,INT_P05\n1,2\n")
        secret = self.pkg / "secret"
        secret.mkdir()
        self.lock(secret)
        with self.assertRaises(phase1_runtime.Phase1Error) as caught:
            phase1_runtime.source_files(self.deposit.root)
        self.assertIn("cannot be listed", str(caught.exception))

    def test_directory_symlinks_are_not_followed(self) -> None:
        (self.pkg / "study.csv").write_text("x,INT_P05\n1,2\n")
        os.symlink(self.pkg, self.pkg / "loop", target_is_directory=True)
        os.symlink(self.pkg / "gone.csv", self.pkg / "dangling.csv")
        vocabulary = self.vocabulary()
        self.assertEqual(vocabulary.unscannable, ())
        self.assertEqual(vocabulary.exempt, frozenset({"INT_P05"}))


class RuleStatementTests(unittest.TestCase):
    def test_statement_names_both_sides_of_the_rule(self) -> None:
        statement = provider_safety.RULE_STATEMENT
        self.assertIn("Never write a token shaped like a participant code", statement)
        self.assertIn("The rule is mechanical", statement)
        self.assertIn("is not an identifier and may be quoted", statement)
        # 2.6.5 review: the validator refuses a source join or evidence path naming such a file,
        # so the statement says so and names the way out.
        self.assertIn("It covers source joins and evidence paths as well", statement)
        self.assertIn("must not be cited in a source join or an evidence path", statement)
        self.assertIn("bind that evidence to another manifest entry", statement)

    def test_the_validator_still_refuses_a_join_to_a_code_named_file(self) -> None:
        sha = hashlib.sha256(b"time,force\n").hexdigest()
        for report, result, rule in (
            (f"## Findings\n<!-- source:pkg/EMG_P01.csv@{sha} -->\n", {}, "report_participant_code"),
            ("## Findings\nok\n", {"scientific_fields": {"findings": [{"evidence": [
                {"source_path": "pkg/EMG_P01.csv", "source_sha256": sha}]}]}}, "result_participant_code"),
        ):
            with self.subTest(rule=rule):
                with self.assertRaises(phase1_runtime.OutputValidationError) as caught:
                    phase1_runtime.check_output_hygiene(report, result, frozenset())
                self.assertEqual(caught.exception.rule, rule)
                self.assertFalse(caught.exception.repairable)


if __name__ == "__main__":
    unittest.main()


class LegacyVerifierHygieneTests(unittest.TestCase):
    """participant_code_rule=False exists only for re-admitting a pre-2.6.5 verifier record."""

    def test_the_rule_can_be_switched_off_only_explicitly(self) -> None:
        report = "## Findings\nThe SUBJ_P12 file is empty.\n"
        with self.assertRaises(phase1_runtime.OutputValidationError) as caught:
            phase1_runtime.check_output_hygiene(report, {})
        self.assertEqual(caught.exception.rule, "report_participant_code")
        phase1_runtime.check_output_hygiene(report, {}, participant_code_rule=False)
        # The other rules still apply to a legacy record.
        with self.assertRaises(phase1_runtime.OutputValidationError):
            phase1_runtime.check_output_hygiene("/Users/someone/data.csv\n", {}, participant_code_rule=False)

