#!/usr/bin/env python3
"""Deterministic, provider-independent support for the packaged Phase-1 runner.

This module deliberately contains no provider launch policy.  It prepares and validates the
immutable inputs on which a provider call may depend, and validates role outputs before they are
eligible for promotion.  Keeping these operations provider-free makes the fan-out fail closed: a
bad source tree, sandbox contract, or promotion path is discovered before any auditor inference.
"""

from __future__ import annotations

import collections
import contextlib
import csv
import fcntl
import hashlib
import json
import os
import pathlib
import posixpath
import platform
import re
import shutil
import stat
import subprocess
import tempfile
import threading
import time
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from typing import Any, Callable, Iterable, Iterator, Union

import sys as _sys

_SCRIPTS_DIR = pathlib.Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in _sys.path:
    _sys.path.insert(0, str(_SCRIPTS_DIR))
import provider_safety  # noqa: E402  -- the participant-code rule shared with run_coordinator
import email_scan  # noqa: E402  -- 2.7.6: linear-time e-mail matching


PHASE1_VERSION = "i4r-phase1-2"
ROLE_RESULT_VERSION = "i4r-phase1-role-result-2"
EVIDENCE_VERSION = "i4r-phase1-evidence-1"
SNAPSHOT_VERSION = "i4r-phase1-snapshot-1"

ROLE_SPECS: dict[str, dict[str, Any]] = {
    "auditor_reproduction": {
        "area": 1,
        "name": "Reproduction",
        "report": "report/01_reproduction.md",
        "work": "work/a1_reproduction",
        "evidence": [
            "pdf_text.txt", "file_inventory.json", "workbook_summaries.json",
            "code_entrypoints.json", "environment_fingerprint.json",
        ],
    },
    "auditor_data_integrity": {
        "area": 2,
        "name": "Data integrity",
        "report": "report/02_data_integrity.md",
        "work": "work/a2_data_integrity",
        "evidence": [
            "pdf_text.txt", "file_inventory.json", "workbook_summaries.json",
            "privacy_scan.json",
        ],
    },
    "auditor_text_vs_data": {
        "area": 3,
        "name": "Text-vs-data consistency",
        "report": "report/03_text_vs_data.md",
        "work": "work/a3_text_vs_data",
        "evidence": [
            "pdf_text.txt", "file_inventory.json", "workbook_summaries.json",
            "code_entrypoints.json",
        ],
    },
    "auditor_code_methods": {
        "area": 4,
        "name": "Code and methods",
        "report": "report/04_code_methods.md",
        "work": "work/a4_code_methods",
        "evidence": [
            "pdf_text.txt", "file_inventory.json", "code_entrypoints.json",
            "environment_fingerprint.json", "archive_licence_metadata.json",
        ],
    },
    "auditor_privacy_archive": {
        "area": 5,
        "name": "Privacy and archive hygiene",
        "report": "report/05_privacy_archive.md",
        "work": "work/a5_privacy_archive",
        "evidence": [
            "file_inventory.json", "workbook_summaries.json", "privacy_scan.json",
            "archive_licence_metadata.json", "environment_fingerprint.json",
        ],
    },
}

# Claude's permission patterns are intentionally data, not prose hidden in a prompt.  The same
# tuple is hashed into the probe key, checked by the local smoke test, and installed in every role
# settings file.  Commands remain confined by the filesystem and network sandbox.
BASH_ALLOWLIST: tuple[str, ...] = (
    # File staging. SKILL.md requires each role to copy code out of the immutable pkg/
    # into its own work/ directory before running it, which needs mkdir/cp/cd/mv. Under
    # --bare the model quietly used the Write tool instead and never exercised this gap;
    # in safe mode it follows the instruction literally and every role was denied.
    # These grant nothing over the Write and Edit tools already in NON_BASH_TOOLS: the
    # real boundary is the filesystem sandbox allowRead/allowWrite plus deny-all network,
    # both unchanged. Redirection stays unsupported by the matcher, so roles capture
    # output from inside their scripts.
    "Bash(mkdir *)",
    "Bash(cp *)",
    "Bash(mv *)",
    "Bash(cd *)",
    "Bash(ls *)",
    "Bash(cat *)",
    "Bash(head *)",
    "Bash(tail *)",
    "Bash(python3 *)",
    "Bash(python *)",
    "Bash(Rscript *)",
    "Bash(R CMD *)",
    "Bash(jupyter *)",
    "Bash(quarto *)",
    "Bash(pdftotext *)",
    "Bash(pdfinfo *)",
    "Bash(file *)",
    "Bash(find *)",
    "Bash(rg *)",
    "Bash(wc *)",
    "Bash(sort *)",
    "Bash(uniq *)",
    "Bash(diff *)",
    "Bash(cmp *)",
    "Bash(shasum *)",
    "Bash(sha256sum *)",
    "Bash(unzip -l *)",
    # Text utilities and cleanup. Claude Code checks every sub-command of a compound
    # command against this list, so with these present a role can stage, run and inspect
    # in one call (`mkdir -p w && cp pkg/a.R w/ && Rscript w/a.R | head`) instead of four.
    # Denials are no longer fatal, but each one still costs a full turn of cache re-read,
    # and turn count is the dominant Phase-1 cost. rm is needed to satisfy the closed
    # output contract; the filesystem sandbox confines it to the attempt output directory.
    "Bash(echo *)",
    "Bash(printf *)",
    "Bash(awk *)",
    "Bash(sed *)",
    "Bash(grep *)",
    "Bash(cut *)",
    "Bash(tr *)",
    "Bash(paste *)",
    "Bash(tee *)",
    "Bash(basename *)",
    "Bash(dirname *)",
    "Bash(pwd)",
    "Bash(test *)",
    "Bash(true)",
    "Bash(touch *)",
    "Bash(rm *)",
    "Bash(rmdir *)",
)
NON_BASH_TOOLS: tuple[str, ...] = ("Read", "Glob", "Grep", "Write", "Edit")

TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".r", ".rmd", ".qmd", ".py", ".do", ".sps", ".sas", ".m", ".tex",
    ".html", ".xml", ".toml", ".ini", ".cfg", ".log",
}
CODE_SUFFIXES = {
    ".r": "R", ".rmd": "R Markdown", ".qmd": "Quarto", ".py": "Python",
    ".do": "Stata", ".ado": "Stata", ".sps": "SPSS", ".sas": "SAS",
    ".ipynb": "Jupyter", ".m": "MATLAB/Octave", ".jl": "Julia",
}
ARCHIVE_SUFFIXES = {".zip", ".tar", ".gz", ".tgz", ".7z", ".rar", ".xz", ".bz2"}
LICENCE_NAMES = re.compile(r"(?i)(?:^|[-_.])(licen[cs]e|copying|notice)(?:[-_.]|$)")
ENTRYPOINT_NAMES = re.compile(
    r"(?i)^(?:main|run|analysis|analyse|master|reproduce|replicate|make|pipeline|workflow)(?:[-_.].*)?$"
)
PARTICIPANT_ID = re.compile(r"(?i)(?<![A-Za-z0-9])(?:participant|subject|subj|pid|sid)[-_ ]?[A-Za-z0-9]{1,16}(?![A-Za-z0-9])")
EMAIL = re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b")
DATE = re.compile(r"\b(?:19|20)\d{2}[-/]\d{1,2}[-/]\d{1,2}\b|\b\d{1,2}:\d{2}(?::\d{2})?\b")
# 2.7.7: the participant/date admission rule (check_output_hygiene) needs an identifier and a real date or clock time.
# PARTICIPANT_ID matches the plain word ("participants", "subject to") and DATE any d:dd, so a sleep study's mean
# bedtime (23:23) or a condition named by a ratio (20:80) on a line that says "participants" was refused as a
# participant linked to a time: one audit was parked before Phase 1 and one role failed all six attempts. Both patterns
# above stay as they are, because the privacy scan counts with them and those counts are Phase-1 evidence.
# An identifier is what follows the label: a token holding a digit ("participant 12", "subject07", "Subj. 12", "#12",
# "sub-01", also in backticks, quotes or bold), a label word and a code ("Participant ID 12", "whose ID is 12",
# "participant_id != 'S07'"), one to four capitals that are not a statistic ("Participant A", "JMK", "TEN-07"; not
# "SD", "RPE", "TST"; "Participant F" unless a sign follows, as in "Participant N = 24"), or a list of codes
# ("participants 12 and 14", "Subjects `S03` and `S07`"). A lowercase word, a range or a count is not one ("participants
# were", "participants 18 to 30 years old", "subjects 1-24", "57 participants"), nor is a capitalised word
# ("Participant Characteristics"), so a named pseudonym ("Participant Anna") is admitted. A time is a clock time
# (hour 0-24, minutes 00-59), also after an ISO date's T, so ratios such as 20:80 and 50:50 are not times.
_ID_SEP = "[-_ \u00a0]"
_ID_Q = "[*`\"'\u2018\u2019\u201c\u201d]{0,2}"   # markdown, backticks and quotes around a code
_ID_CODE = r"#?(?:[A-Za-z]{1,8}[-_.])?(?=[A-Za-z0-9]{0,15}\d)[A-Za-z0-9]{1,16}(?![A-Za-z0-9])"
_ID_ITEM = rf"{_ID_Q}#?(?:[A-Za-z]{{1,4}}[-_.]?)?\d{{1,4}}[A-Za-z]?(?![A-Za-z0-9]){_ID_Q}"
_ID_LABEL = r"(?:ids?|identifiers?|nos?|nr|n°|nº|num|numbers?|codes?)\.?"
_NOT_A_PSEUDONYM = (r"(?:SDS?|SEM?|CVS?|CIS?|ICC|IQR|IDS?|AM|PM|RTS?|MD|N|M|F|OR|AND|THE|OF|TO|IN|RPE|HRV?|BMI|BP|MVC|"
                    r"EMG|EEG|ECG|MRI|CMJ|ACWR|RSI|SRPE|TL|VAS|PSQI|ESS|MEQ|POMS|TST|SOL|WASO)")
PARTICIPANT_LABELLED_ID = re.compile(   # every separator run is bounded: unbounded adjacent runs are quadratic
    r"(?i)(?<![A-Za-z0-9])(?:(?:participant|subject|subj|pid|sid)(?:"
    rf"\.?{_ID_SEP}?{_ID_Q}{_ID_CODE}"
    # a label word, after "with"/"whose"/"of" (not "with no 3 valid nights"), and an operator, copula or quote
    rf"|s?(?:{_ID_SEP}{{1,3}}(?:with|whose|of)(?:{_ID_SEP}{{1,3}}the)?(?!{_ID_SEP}{{1,3}}(?-i:no)(?![.A-Za-z0-9])))?"
    rf"{_ID_SEP}{{0,3}}{_ID_LABEL}(?:{_ID_SEP}{{1,3}}(?:no|number)\.?)?(?:{_ID_SEP}{{1,3}}(?:is|was|are|were))?"
    rf"[-_ \u00a0:#=!<>\x22']{{0,6}}{_ID_Q}{_ID_CODE}"
    rf"|{_ID_SEP}(?-i:(?!{_NOT_A_PSEUDONYM}(?![A-Za-z0-9]))[A-Z]{{1,4}})(?![A-Za-z0-9])"
    rf"|s?{_ID_SEP}{{0,3}}:?{_ID_SEP}{{0,3}}{_ID_ITEM}(?:{_ID_SEP}{{0,3}}(?:,|and|&){_ID_SEP}{{0,3}}{_ID_ITEM})+)"
    # a single M, N or F after a capitalised label is a pseudonym unless a statistic's sign follows ("Participant N = 24")
    rf"|(?-i:Participant|Subject|Subj\.?){_ID_SEP}(?-i:[MNF])(?![A-Za-z0-9])(?!{_ID_SEP}{{0,3}}[=<>(\u2248]))")
DATE_OR_CLOCK_TIME = re.compile(
    r"\b(?:19|20)\d{2}[-/]\d{1,2}[-/]\d{1,2}(?!\d)"
    r"|\b(?:[01０１]?\d|[2２][0-4０-４]):[0-5０-５]\d(?::[0-5０-５]\d)?\b"
    r"|(?<=\d[Tt])(?:[01]\d|2[0-3]):[0-5]\d(?::[0-5]\d)?(?!\d)")

def participant_date_linked(line: str) -> bool:
    """The participant_date admission rule (2.7.7). It also requires the 2.7.6 pattern pair, so it only ever admits
    more than 2.7.6 did: a report admitted before, such as a verifier record re-admitted later, is admitted again."""
    return bool(PARTICIPANT_ID.search(line) and DATE.search(line)
                and PARTICIPANT_LABELLED_ID.search(line) and DATE_OR_CLOCK_TIME.search(line))


ABSOLUTE_PATH = re.compile(
    r"(?:/" + "Users/" + r"|/home/|/private/(?:tmp|var)/|[A-Za-z]:[\\/]Users[\\/])"
)
REPORT_SOURCE_BINDING = re.compile(
    r"<!--\s*source:([^@\s][^@]*?)@([0-9a-f]{64})\s*-->", re.IGNORECASE
)
PLACEHOLDER = re.compile(r"(?im)^\s*(?:todo|tbd|placeholder|<insert[^>]*>|\[insert[^]]*\])\s*$")


class Phase1Error(ValueError):
    """Expected fail-closed Phase-1 error."""


class SnapshotIntegrityError(Phase1Error):
    """The frozen Phase-1 snapshot no longer matches its manifest: an environment fault.

    Raised for content, manifest and structure failures (and for mode drift that will not stay
    re-frozen). No model output can cause it and a fresh sample cannot cure it, so the runners
    pause on it instead of spending attempts. Messages name manifest indexes, never deposit
    file names: they are recorded in invocations.jsonl and projected into run.json.
    """


class SnapshotModeUnstableError(SnapshotIntegrityError):
    """The snapshot's content verified, but its modes would not stay frozen (2.6.8).

    Raised after a bounded number of re-freezes when something on the host keeps giving a frozen
    entry its write bit back. It is still a SnapshotIntegrityError (an environment fault, never a
    model's), but it is not a content change: deleting the snapshot cannot help while the modes
    keep changing, so its message and pause reason (``snapshot_mode_unstable``) say to stop the sync
    client or move the audit off the synced folder instead.
    """


SNAPSHOT_MODE_UNSTABLE = "snapshot_mode_unstable"
# A re-freeze is repeated this many times, with a short growing pause, before a mode that keeps
# changing is reported as SnapshotModeUnstableError: one File Provider reset is repaired quietly.
SNAPSHOT_REFREEZE_ATTEMPTS = 3
SNAPSHOT_REFREEZE_RETRY_SECONDS = 0.1


def snapshot_delete_remedy(relative: str) -> str:
    """The one command that removes a frozen snapshot: 0555 directories refuse a plain rm -rf."""
    return f"chmod -R u+w {relative} && rm -rf {relative}"


class OutputValidationError(Phase1Error):
    """A provider output failed the local admission contract.

    ``rule`` names the admission rule that rejected the output. The auditor path never reads it;
    the 2.6.4 verifier's prompt-contract test uses it to prove that every rule its validator can
    reject on is stated in the verifier prompt.
    """

    def __init__(self, message: str, *, repairable: bool = False, rule: str | None = None) -> None:
        super().__init__(message)
        self.repairable = repairable
        self.rule = rule


@contextlib.contextmanager
def _admission_rule(rule: str) -> Iterator[None]:
    """Attribute a helper's bare Phase1Error (a missing file, bad JSON, an escaping path) to one
    admission rule. The message is unchanged, and validate_role_output already turned these into
    non-repairable OutputValidationErrors with the same text, so auditor behaviour is identical."""
    try:
        yield
    except OutputValidationError as exc:
        if exc.rule is None:
            exc.rule = rule
        raise
    except Phase1Error as exc:
        raise OutputValidationError(str(exc), rule=rule) from exc


@dataclass(frozen=True)
class MonotonicDeadline:
    """One shared wall-clock budget for every provider-free preparation operation."""

    expires_at: float

    @classmethod
    def after(cls, seconds: float) -> "MonotonicDeadline":
        if isinstance(seconds, bool) or not isinstance(seconds, (int, float)) or seconds <= 0:
            raise Phase1Error("pre-dispatch deadline must be a positive number of seconds")
        return cls(time.monotonic() + float(seconds))

    def remaining(self, label: str, *, cap: float | None = None) -> float:
        remaining = self.expires_at - time.monotonic()
        if remaining <= 0:
            raise Phase1Error(f"Phase-1 pre-dispatch deadline exceeded during {label}")
        return min(remaining, cap) if cap is not None else remaining

    def check(self, label: str) -> None:
        self.remaining(label)

    def subprocess_timeout(self, label: str, *, cap: float) -> float:
        """Leave a small cleanup margin before the enclosing hard deadline fires."""
        remaining = self.remaining(label, cap=cap)
        return max(0.001, remaining - min(0.1, remaining / 10))


def check_deadline(deadline: MonotonicDeadline | None, label: str) -> None:
    if deadline is not None:
        deadline.check(label)


@dataclass(frozen=True)
class ValidatedRoleOutput:
    role: str
    report: pathlib.Path
    evidence_manifest: pathlib.Path
    result: pathlib.Path
    report_sha256: str
    evidence_manifest_sha256: str
    result_sha256: str
    scientific_fields: dict[str, Any]
    source_bindings: tuple[tuple[str, str], ...]


def canonical_bytes(value: Any) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode("utf-8")


def pretty_bytes(value: Any) -> bytes:
    return (json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n").encode("utf-8")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: pathlib.Path, deadline: MonotonicDeadline | None = None) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            check_deadline(deadline, f"hashing {path.name}")
            block = handle.read(1024 * 1024)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def atomic_write(path: pathlib.Path, data: bytes, mode: int = 0o600) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    fd = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, mode)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def strict_json(path: pathlib.Path) -> dict[str, Any]:
    def unique_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        value: dict[str, Any] = {}
        for key, item in pairs:
            if key in value:
                raise ValueError(f"duplicate object key {key!r}")
            value[key] = item
        return value

    try:
        value = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=unique_object)
    except (OSError, UnicodeError, json.JSONDecodeError, ValueError) as exc:
        raise Phase1Error(f"invalid JSON object {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise Phase1Error(f"expected JSON object: {path}")
    return value


def regular_file(path: pathlib.Path, label: str) -> pathlib.Path:
    if path.is_symlink() or not path.is_file():
        raise Phase1Error(f"{label} must be a regular non-symlink file: {path}")
    return path


def beneath(root: pathlib.Path, path: pathlib.Path, label: str) -> pathlib.Path:
    resolved = path.resolve()
    try:
        resolved.relative_to(root.resolve())
    except ValueError as exc:
        raise Phase1Error(f"{label} escapes {root}: {path}") from exc
    return resolved


def validate_destination(root: pathlib.Path, path: pathlib.Path, label: str) -> None:
    """Reject a destination whose existing component is a symlink or escapes the audit."""
    beneath(root, path, label)
    current = path
    while current != root:
        if current.is_symlink():
            raise Phase1Error(f"{label} has a symlink component: {current}")
        current = current.parent
    if path.exists() and not path.is_file() and not path.is_dir():
        raise Phase1Error(f"{label} is a special file: {path}")


def source_files(
    root: pathlib.Path, deadline: MonotonicDeadline | None = None,
) -> list[tuple[str, pathlib.Path]]:
    """Return every raw paper/package source, rejecting symlinks and special files."""
    rows: list[tuple[str, pathlib.Path]] = []
    for base_name in ("paper", "pkg"):
        check_deadline(deadline, f"traversing {base_name}")
        base = root / base_name
        if base.is_symlink() or not base.is_dir():
            raise Phase1Error(f"required source directory is missing or unsafe: {base}")

        # rglob skips a directory it cannot list without a word. The coordinator's deposit scan
        # (provider_safety, 2.6.5) treats such a deposit as unscannable and exempts nothing, so a
        # snapshot that silently dropped the directory would let Phase 1 admit, and pay for,
        # reports the coordinator then refuses. Refuse it here, before any provider call.
        def unlistable(exc: OSError) -> None:
            raise Phase1Error(
                f"sandbox source directory cannot be listed: {exc.filename}: {exc.strerror or exc}"
            )

        for _current, _dirs, _files in os.walk(base, onerror=unlistable):
            check_deadline(deadline, f"traversing {base_name}")
        discovered: list[pathlib.Path] = []
        for path in base.rglob("*"):
            check_deadline(deadline, f"traversing {base_name}")
            discovered.append(path)
        for path in sorted(discovered, key=lambda item: item.as_posix()):
            check_deadline(deadline, f"traversing {base_name}")
            if path.is_symlink():
                # Three cases, and only one is a sandbox problem.
                #   escaping  -> fatal: reading it leaves the frozen tree.
                #   dangling  -> skip: it has no bytes to hash or read. One deposit ships
                #                conf.yaml -> src/conf.yaml with no src/ in the archive; that is
                #                broken packaging by the depositor, not a breach, and refusing it
                #                stranded the audit after the paper was already reserved.
                #   internal  -> follow it and treat the target as an ordinary source file.
                try:
                    target = path.resolve(strict=True)
                except (OSError, RuntimeError):
                    continue
                try:
                    target.relative_to(base.resolve())
                except ValueError:
                    raise Phase1Error(
                        f"sandbox source contains a symlink escaping {base_name}: {path}"
                    ) from None
                if target.is_dir():
                    continue
            if path.is_dir():
                continue
            if not path.is_file():
                raise Phase1Error(f"sandbox source contains a special file: {path}")
            try:
                with path.open("rb") as handle:
                    handle.read(1)
            except OSError as exc:
                raise Phase1Error(f"sandbox source is not readable: {path}: {exc}") from exc
            rows.append((path.relative_to(root).as_posix(), path))
    if not any(relative.startswith("paper/") for relative, _ in rows):
        raise Phase1Error("paper/ contains no readable source file")
    return rows


def verify_package_manifest(
    root: pathlib.Path, files: list[tuple[str, pathlib.Path]],
    deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    manifest_path = regular_file(root / "package_manifest.json", "package manifest")
    manifest = strict_json(manifest_path)
    declared = manifest.get("files")
    if not isinstance(declared, list):
        raise Phase1Error("package_manifest.json files must be an array")
    expected: dict[str, tuple[str, int]] = {}
    for index, row in enumerate(declared):
        check_deadline(deadline, "validating package manifest")
        if not isinstance(row, dict) or not isinstance(row.get("path"), str):
            raise Phase1Error(f"package manifest file row {index} is invalid")
        path = row["path"]
        digest = row.get("sha256")
        size = row.get("size", row.get("bytes"))
        if not re.fullmatch(r"[0-9a-f]{64}", str(digest)) or not isinstance(size, int) or size < 0:
            raise Phase1Error(f"package manifest file row {index} lacks exact hash/size")
        if path in expected:
            raise Phase1Error(f"duplicate package manifest path: {path}")
        expected[path] = (str(digest), size)
    actual_paths = {relative[4:]: path for relative, path in files if relative.startswith("pkg/")}
    if set(expected) != set(actual_paths):
        missing = sorted(set(expected) - set(actual_paths))[:5]
        extra = sorted(set(actual_paths) - set(expected))[:5]
        raise Phase1Error(f"package manifest closure mismatch; missing={missing}, extra={extra}")
    for relative, (digest, size) in expected.items():
        check_deadline(deadline, "validating package hashes")
        path = actual_paths[relative]
        if path.stat().st_size != size or sha256_file(path, deadline) != digest:
            raise Phase1Error(f"package manifest hash/size mismatch: pkg/{relative}")
    declared_tree = manifest.get("package_sha256")
    if declared_tree is not None:
        tree = hashlib.sha256(
            # the manifest's own order: pathlib's, part by part (setup_workspace, init_run and the post-delta gates
            # all use it; a plain string sort puts "code.R" before "code/x.R" and reads as a tampered manifest)
            "".join(f"{path}\0{expected[path][0]}\n"
                    for path in sorted(expected, key=lambda value: pathlib.PurePosixPath(value).parts)).encode("utf-8")
        ).hexdigest()
        if declared_tree != tree:
            raise Phase1Error("package manifest aggregate package_sha256 is invalid")
    return manifest


def file_inventory(
    files: list[tuple[str, pathlib.Path]], deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    types: collections.Counter[str] = collections.Counter()
    rows: list[dict[str, Any]] = []
    for relative, path in files:
        check_deadline(deadline, "building file inventory")
        suffix = path.suffix.lower().lstrip(".") or "<none>"
        types[suffix] += 1
        rows.append({
            "path": relative,
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path, deadline),
            "extension": suffix,
        })
    return {
        "schema_version": EVIDENCE_VERSION,
        "file_count": len(rows),
        "file_types": dict(sorted(types.items())),
        "files": rows,
    }


def _csv_rows(path: pathlib.Path, delimiter: str) -> Iterator[list[str]]:
    with path.open("r", encoding="utf-8-sig", errors="strict", newline="") as handle:
        yield from csv.reader(handle, delimiter=delimiter)


def _summarize_rows(
    rows: Iterable[Iterable[Any]], *, sheet_index: int, sheet_name: str,
    deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    row_count = 0
    max_columns = 0
    missing_by_column: list[int] = []
    types: collections.Counter[str] = collections.Counter()
    row_digests: collections.Counter[str] = collections.Counter()
    for raw in rows:
        check_deadline(deadline, f"scanning workbook sheet {sheet_name}")
        values = list(raw)
        row_count += 1
        max_columns = max(max_columns, len(values))
        while len(missing_by_column) < len(values):
            missing_by_column.append(row_count - 1)
        normalized: list[Any] = []
        for index in range(max_columns):
            value = values[index] if index < len(values) else None
            missing = value is None or (isinstance(value, str) and not value.strip())
            if missing:
                if index >= len(missing_by_column):
                    missing_by_column.append(row_count)
                else:
                    missing_by_column[index] += 1
                normalized.append(None)
                types["missing"] += 1
            elif isinstance(value, bool):
                normalized.append(value)
                types["boolean"] += 1
            elif isinstance(value, (int, float)):
                normalized.append(value)
                types["number"] += 1
            else:
                text = str(value)
                normalized.append(text)
                try:
                    float(text)
                except ValueError:
                    types["text"] += 1
                else:
                    types["number_text"] += 1
        row_digests[sha256_bytes(canonical_bytes(normalized))] += 1
    duplicate_rows = sum(count - 1 for count in row_digests.values() if count > 1)
    return {
        "sheet_index": sheet_index,
        "sheet_name": sheet_name,
        "row_count": row_count,
        "column_count": max_columns,
        "missing_by_column_index": missing_by_column,
        "cell_types": dict(sorted(types.items())),
        "duplicate_row_count": duplicate_rows,
    }


# Measured against this tree, not guessed. Row summarisation costs roughly 60-100 ms per MB of
# CSV: the two files above this line cost 21.1 s (334 MB) and 29.3 s (374 MB) out of a 60-second
# pre-dispatch deadline shared with everything else, while the largest workbook below it -- a
# 20.3 MB xlsx with 387k rows -- costs 6.1 s and is exactly the kind of dataset a reproduction
# check needs. The first cut at this was 20 MB, which bought little and dropped that workbook.
WORKBOOK_SUMMARY_MAX_BYTES = 256 * 1024 * 1024

# Size is only a proxy for cost, so the real bound is time. A single file may spend this long
# before it is abandoned as unavailable, and the whole scan may spend the second figure; past
# either, the remaining files are inventoried and hashed but not summarised. Without these the
# only protection was the shared deadline, and blowing that fails the entire preflight -- after
# the paper is reserved and pkg/ frozen -- instead of degrading one file's evidence.
WORKBOOK_SUMMARY_FILE_SECONDS = 8.0
WORKBOOK_SUMMARY_TOTAL_SECONDS = 20.0

# Suffixes workbook_summaries can actually summarise, plus the two it reports as unavailable
# because no deterministic reader is installed. Anything else is not a workbook at all.
WORKBOOK_SUFFIXES = frozenset({".csv", ".tsv", ".xlsx", ".xls", ".ods"})


def workbook_summaries(
    files: list[tuple[str, pathlib.Path]], deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    books: list[dict[str, Any]] = []
    scan_deadline = MonotonicDeadline(time.monotonic() + WORKBOOK_SUMMARY_TOTAL_SECONDS)
    for relative, path in files:
        check_deadline(deadline, "scanning workbooks")
        suffix = path.suffix.lower()
        # The suffix test comes FIRST. It used to sit below the size cap, as the `else: continue`
        # at the end of the dispatch, so every oversize file of any kind was appended as an
        # "oversize workbook" -- 34 of them in this tree (.zip, .rds, .mov, .sav) across 22
        # papers, each inflating workbook_count and telling the role to "read the file directly"
        # about a movie. None of them would have been summarised even at one byte.
        if suffix not in WORKBOOK_SUFFIXES:
            continue
        summaries: list[dict[str, Any]] = []
        status = "summarized"
        reason: str | None = None
        # Row-by-row summarisation has no natural bound: a 1.19 GB CSV in one deposit takes 178 s
        # against a 60-second pre-dispatch deadline that cannot be raised, so the whole audit
        # pauses at preflight after the paper is already reserved and pkg/ already frozen. The
        # file is still inventoried and hashed; only the deterministic row summary is skipped,
        # and the role is told so rather than guessing.
        size = path.stat().st_size
        if size > WORKBOOK_SUMMARY_MAX_BYTES:
            books.append({
                "path": relative, "sha256": sha256_file(path, deadline),
                "format": suffix.lstrip("."), "status": "unavailable",
                "unavailable_reason": (
                    "oversize for a deterministic summary "
                    f"({size} bytes > {WORKBOOK_SUMMARY_MAX_BYTES}); "
                    "read the file directly if a claim depends on it"),
                "sheets": [],
            })
            continue
        file_deadline = MonotonicDeadline(min(
            scan_deadline.expires_at, time.monotonic() + WORKBOOK_SUMMARY_FILE_SECONDS,
            *( (deadline.expires_at,) if deadline is not None else () ),
        ))
        try:
            if suffix == ".csv":
                summaries.append(_summarize_rows(
                    _csv_rows(path, ","), sheet_index=1, sheet_name="CSV", deadline=file_deadline,
                ))
            elif suffix == ".tsv":
                summaries.append(_summarize_rows(
                    _csv_rows(path, "\t"), sheet_index=1, sheet_name="TSV", deadline=file_deadline,
                ))
            elif suffix == ".xlsx":
                try:
                    import openpyxl  # type: ignore
                except ImportError:
                    status, reason = "unavailable", "openpyxl is not installed"
                else:
                    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
                    try:
                        for index, sheet in enumerate(workbook.worksheets, 1):
                            summaries.append(_summarize_rows(
                                sheet.iter_rows(values_only=True), sheet_index=index,
                                sheet_name=sheet.title, deadline=file_deadline,
                            ))
                    finally:
                        workbook.close()
            else:  # .xls / .ods -- inventoried and hashed, but no deterministic reader here
                status, reason = "unavailable", f"deterministic reader for {suffix} is not installed"
        except Phase1Error:
            # Either the per-file/scan budget or the shared pre-dispatch deadline. Re-check the
            # shared one: if it is genuinely gone the whole preflight must still fail, and only
            # a budget of our own degrades to unavailable evidence for this one file.
            check_deadline(deadline, "scanning workbooks")
            status, summaries = "unavailable", []
            reason = (
                f"deterministic summary exceeded its {WORKBOOK_SUMMARY_FILE_SECONDS:g}s per-file "
                f"budget (or the {WORKBOOK_SUMMARY_TOTAL_SECONDS:g}s budget for the whole "
                "workbook scan); read the file directly if a claim depends on it")
        except (OSError, UnicodeError, csv.Error, ValueError, zipfile.BadZipFile) as exc:
            status, reason, summaries = "unavailable", f"{type(exc).__name__}: {exc}", []
        books.append({
            "path": relative,
            "sha256": sha256_file(path, deadline),
            "format": suffix.lstrip("."),
            "status": status,
            "unavailable_reason": reason,
            "sheets": summaries,
        })
    return {"schema_version": EVIDENCE_VERSION, "workbook_count": len(books), "workbooks": books}


def code_entrypoints(
    files: list[tuple[str, pathlib.Path]], deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for relative, path in files:
        check_deadline(deadline, "indexing code entrypoints")
        language = CODE_SUFFIXES.get(path.suffix.lower())
        if not language:
            continue
        reasons: list[str] = []
        if ENTRYPOINT_NAMES.match(path.stem):
            reasons.append("entrypoint_filename")
        if path.suffix.lower() == ".ipynb":
            reasons.append("notebook")
        if path.suffix.lower() in {".rmd", ".qmd"}:
            reasons.append("literate_analysis")
        if path.stat().st_size <= 2 * 1024 * 1024 and path.suffix.lower() in TEXT_SUFFIXES:
            try:
                text = path.read_text(encoding="utf-8", errors="strict")
            except (OSError, UnicodeError):
                text = ""
            if re.search(r"(?m)^\s*if\s+__name__\s*==\s*['\"]__main__['\"]", text):
                reasons.append("python_main_guard")
            if re.search(r"(?im)^\s*(?:source|run|do)\s*\(", text):
                reasons.append("orchestrates_other_files")
        rows.append({
            "path": relative,
            "sha256": sha256_file(path, deadline),
            "language": language,
            "entrypoint_signals": sorted(set(reasons)),
        })
    return {"schema_version": EVIDENCE_VERSION, "code_file_count": len(rows), "files": rows}


def _tool_version(
    command: list[str], deadline: MonotonicDeadline | None = None,
) -> str | None:
    check_deadline(deadline, f"probing {command[0]} version")
    executable = shutil.which(command[0])
    if not executable:
        return None
    try:
        result = subprocess.run(
            [executable, *command[1:]], text=True, stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=deadline.subprocess_timeout(f"probing {command[0]} version", cap=10)
            if deadline is not None else 10,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    first = next((line.strip() for line in result.stdout.splitlines() if line.strip()), "")
    # Tool versions occasionally include a home path; do not preserve it in common evidence.
    return ABSOLUTE_PATH.sub("<private-path>/", first)[:300] or None


def environment_fingerprint(
    deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    check_deadline(deadline, "building environment fingerprint")
    value = {
        "schema_version": EVIDENCE_VERSION,
        "platform_system": platform.system(),
        "platform_machine": platform.machine(),
        "python_implementation": platform.python_implementation(),
        "python_version": platform.python_version(),
        "tools": {
            "R": _tool_version(["R", "--version"], deadline),
            "python": _tool_version(["python3", "--version"], deadline),
            "quarto": _tool_version(["quarto", "--version"], deadline),
            "stata_executable_present": bool(shutil.which("stata") or shutil.which("stata-mp")),
            "pdftotext": _tool_version(["pdftotext", "-v"], deadline),
        },
    }
    value["fingerprint_sha256"] = sha256_bytes(canonical_bytes(value))
    return value


def archive_licence_metadata(
    files: list[tuple[str, pathlib.Path]], deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    archives: list[dict[str, Any]] = []
    licences: list[dict[str, Any]] = []
    for relative, path in files:
        check_deadline(deadline, "indexing archives and licences")
        row = {
            "path": relative, "sha256": sha256_file(path, deadline),
            "bytes": path.stat().st_size,
        }
        if path.suffix.lower() in ARCHIVE_SUFFIXES:
            archives.append({**row, "format": path.suffix.lower().lstrip(".")})
        if LICENCE_NAMES.search(path.name):
            licences.append(row)
    return {
        "schema_version": EVIDENCE_VERSION,
        "archives": archives,
        "licence_or_notice_files": licences,
    }


def privacy_scan(
    files: list[tuple[str, pathlib.Path]], deadline: MonotonicDeadline | None = None,
) -> dict[str, Any]:
    counts = collections.Counter({
        "participant_identifier": 0,
        "email": 0,
        "participant_identifier_near_date_or_time": 0,
        "private_absolute_path": 0,
    })
    scanned = 0
    skipped_binary = 0
    for _relative, path in files:
        check_deadline(deadline, "running privacy scan")
        if path.suffix.lower() not in TEXT_SUFFIXES or path.stat().st_size > 20 * 1024 * 1024:
            skipped_binary += 1
            continue
        try:
            text = path.read_text(encoding="utf-8", errors="strict")
        except (OSError, UnicodeError):
            skipped_binary += 1
            continue
        scanned += 1
        counts["participant_identifier"] += len(PARTICIPANT_ID.findall(text))
        counts["email"] += len(email_scan.findall(EMAIL, text))   # 2.7.6: linear time
        counts["private_absolute_path"] += len(ABSOLUTE_PATH.findall(text))
        for line in text.splitlines():
            check_deadline(deadline, "running privacy scan")
            if PARTICIPANT_ID.search(line) and DATE.search(line):
                counts["participant_identifier_near_date_or_time"] += 1
    return {
        "schema_version": EVIDENCE_VERSION,
        "files_scanned_as_text": scanned,
        "files_skipped_as_binary_or_oversize": skipped_binary,
        "detector_match_counts": dict(sorted(counts.items())),
        "matches_emitted": False,
        "participant_rows_emitted": False,
        "self_tests": {
            "known_present": bool(PARTICIPANT_ID.search("participant P7") and EMAIL.search("x@example.org")),
            "known_absent": not bool(PARTICIPANT_ID.search("aggregate result only") or EMAIL.search("aggregate result only")),
        },
    }


def paper_text(
    root: pathlib.Path, deadline: MonotonicDeadline | None = None,
) -> tuple[bytes, str]:
    candidates = (root / "paper/paper.txt", root / "PAPER.txt")
    for candidate in candidates:
        check_deadline(deadline, "reading paper text")
        if candidate.is_file() and not candidate.is_symlink():
            data = candidate.read_bytes()
            data.decode("utf-8")
            return data, candidate.relative_to(root).as_posix()
    pdf = root / "paper/paper.pdf"
    regular_file(pdf, "paper PDF")
    executable = shutil.which("pdftotext")
    if not executable:
        raise Phase1Error("paper text is missing and pdftotext is unavailable")
    with tempfile.TemporaryDirectory(prefix="i4r-pdf-text-") as directory:
        destination = pathlib.Path(directory) / "paper.txt"
        try:
            result = subprocess.run(
                [executable, "-layout", str(pdf), str(destination)], text=True,
                stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                timeout=deadline.subprocess_timeout("extracting paper PDF", cap=120)
                if deadline is not None else 120,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise Phase1Error("paper extraction exceeded the pre-dispatch deadline") from exc
        check_deadline(deadline, "extracting paper PDF")
        if result.returncode or not destination.is_file():
            raise Phase1Error(f"pdftotext failed: {result.stderr.strip()[:500]}")
        data = destination.read_bytes()
        data.decode("utf-8")
        return data, "generated:pdftotext-layout"


def build_common_evidence(
    root: pathlib.Path, destination: pathlib.Path,
    deadline: MonotonicDeadline | None = None,
) -> tuple[dict[str, Any], list[tuple[str, pathlib.Path]]]:
    files = source_files(root, deadline)
    verify_package_manifest(root, files, deadline)
    text, text_origin = paper_text(root, deadline)
    payloads: dict[str, bytes] = {
        "pdf_text.txt": text,
        "file_inventory.json": pretty_bytes(file_inventory(files, deadline)),
        "workbook_summaries.json": pretty_bytes(workbook_summaries(files, deadline)),
        "code_entrypoints.json": pretty_bytes(code_entrypoints(files, deadline)),
        "environment_fingerprint.json": pretty_bytes(environment_fingerprint(deadline)),
        "archive_licence_metadata.json": pretty_bytes(archive_licence_metadata(files, deadline)),
        "privacy_scan.json": pretty_bytes(privacy_scan(files, deadline)),
    }
    destination.mkdir(parents=True, exist_ok=False)
    artifacts = []
    for name in sorted(payloads):
        check_deadline(deadline, "writing common evidence")
        path = destination / name
        atomic_write(path, payloads[name])
        artifacts.append({
            "path": name, "sha256": sha256_file(path, deadline),
            "bytes": path.stat().st_size,
        })
    index = {
        "schema_version": EVIDENCE_VERSION,
        "generator_version": PHASE1_VERSION,
        "participant_rows_emitted": False,
        "model_conclusions_emitted": False,
        "pdf_text_origin": text_origin,
        "artifacts": artifacts,
    }
    atomic_write(destination / "evidence_index.json", pretty_bytes(index))
    return index, files


def _copy_regular(
    source: pathlib.Path, destination: pathlib.Path,
    deadline: MonotonicDeadline | None = None,
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with source.open("rb") as reader, destination.open("xb") as writer:
        while True:
            check_deadline(deadline, f"copying {source.name}")
            block = reader.read(1024 * 1024)
            if not block:
                break
            writer.write(block)
        writer.flush()
        os.fsync(writer.fileno())
    os.chmod(destination, 0o444)


def freeze_tree(path: pathlib.Path, deadline: MonotonicDeadline | None = None) -> None:
    for directory, subdirs, files in os.walk(path, topdown=False):
        check_deadline(deadline, "freezing source snapshot")
        for name in files:
            os.chmod(pathlib.Path(directory) / name, 0o444)
        for name in subdirs:
            os.chmod(pathlib.Path(directory) / name, 0o555)
        os.chmod(directory, 0o555)


def freeze_and_publish(
    temporary: pathlib.Path, target: pathlib.Path,
    deadline: MonotonicDeadline | None = None,
) -> None:
    """Publish an immutable tree so that no observer can ever see it unfrozen.

    Freezing after os.replace left a published snapshot whose root was still writable if the
    process died in between -- and freeze_tree walks bottom-up, chmod'ing the root last, while
    checking the deadline per directory, so a routine deadline expiry was enough. Every later run
    then failed on "immutable path regained write permission" and needed a manual chmod.

    Freezing first is not sufficient on its own: renaming a directory ACROSS parents updates its
    `..` entry and therefore needs write permission on the directory being moved, so a fully
    frozen tree cannot be moved out of a staging directory. The critique helper avoids this by
    staging inside the destination's own parent, where `..` never changes. Do the same here: move
    the still-writable tree next to its target first, freeze it there, then rename within that one
    directory. Only the hidden .tmp name is ever observable unfrozen, and it is never consumed.
    """
    target.parent.mkdir(parents=True, exist_ok=True)
    staged = target.with_name(f".{target.name}.{os.getpid()}.tmp")
    if staged.exists():
        shutil.rmtree(staged, ignore_errors=True)
    os.replace(temporary, staged)
    try:
        freeze_tree(staged, deadline)
        os.replace(staged, target)
    except BaseException:
        for path in [staged, *staged.rglob("*")] if staged.exists() else []:
            try:
                path.chmod(0o700 if path.is_dir() else 0o600)
            except OSError:
                pass
        shutil.rmtree(staged, ignore_errors=True)
        raise


def materialize_snapshot(
    root: pathlib.Path, phase_root: pathlib.Path,
    deadline: MonotonicDeadline | None = None,
    *,
    refreeze: bool = False,
    on_refreeze: Callable[[dict[str, Any]], None] | None = None,
    lenient_modes: bool = False,
) -> tuple[pathlib.Path, dict[str, Any]]:
    """Build common evidence once and promote one immutable, content-addressed source snapshot.

    ``refreeze``, ``on_refreeze`` and ``lenient_modes`` are passed to verify_snapshot. Inside the
    SIGALRM pre-dispatch deadline pass ``lenient_modes`` and never ``refreeze``: the chmod walk
    must not be interruptible, so the runners re-freeze after the deadline is disarmed.
    """
    staging_parent = phase_root / "snapshot_staging"
    validate_destination(root, phase_root, "Phase-1 runtime directory")
    validate_destination(root, staging_parent, "snapshot staging directory")
    staging_parent.mkdir(parents=True, exist_ok=True)
    temporary = pathlib.Path(tempfile.mkdtemp(prefix="snapshot-", dir=staging_parent))
    # mkdtemp makes it 0700 and freeze_and_publish moves it across parents; the File Provider has
    # been seen re-applying exactly that creation mode to a frozen tree. freeze_tree still sets
    # the frozen mode, and the snapshot hash covers no mode.
    temporary.chmod(0o755)
    try:
        check_deadline(deadline, "starting source snapshot")
        evidence_dir = temporary / "evidence"
        evidence_index, files = build_common_evidence(root, evidence_dir, deadline)
        source_rows = []
        for relative, source in files:
            check_deadline(deadline, "copying source snapshot")
            destination = temporary / "raw" / relative
            _copy_regular(source, destination, deadline)
            source_rows.append({
                "path": relative,
                "sha256": sha256_file(destination, deadline),
                "bytes": destination.stat().st_size,
            })
        package_manifest = root / "package_manifest.json"
        _copy_regular(package_manifest, temporary / "context/package_manifest.json", deadline)
        policy = regular_file(root / "pipeline-policy.snapshot.json", "policy snapshot")
        _copy_regular(policy, temporary / "context/pipeline-policy.snapshot.json", deadline)
        manifest = {
            "schema_version": SNAPSHOT_VERSION,
            "source_files": source_rows,
            "common_evidence": evidence_index,
            "common_evidence_index_sha256": sha256_file(
                evidence_dir / "evidence_index.json", deadline,
            ),
            "context": [
                {
                    "path": "package_manifest.json",
                    "sha256": sha256_file(package_manifest, deadline),
                },
                {
                    "path": "pipeline-policy.snapshot.json",
                    "sha256": sha256_file(policy, deadline),
                },
            ],
        }
        snapshot_hash = sha256_bytes(canonical_bytes(manifest))
        manifest["snapshot_sha256"] = snapshot_hash
        atomic_write(temporary / "snapshot_manifest.json", pretty_bytes(manifest))
        snapshots = phase_root / "snapshots"
        validate_destination(root, snapshots, "snapshot store")
        snapshots.mkdir(parents=True, exist_ok=True)
        target = snapshots / snapshot_hash
        validate_destination(root, target, "content-addressed snapshot")
        if target.exists():
            existing = strict_json(regular_file(target / "snapshot_manifest.json", "snapshot manifest"))
            if existing != manifest:
                raise Phase1Error(f"content-addressed snapshot collision: {target}")
            shutil.rmtree(temporary)
            _verify_published(target, manifest, deadline, refreeze, on_refreeze, lenient_modes)
            return target, manifest
        freeze_and_publish(temporary, target, deadline)
        _verify_published(target, manifest, deadline, refreeze, on_refreeze, lenient_modes)
        return target, manifest
    finally:
        if temporary.exists():
            shutil.rmtree(temporary, ignore_errors=True)


def refreeze_published_snapshot(
    snapshot: pathlib.Path, manifest: dict[str, Any],
    on_refreeze: Callable[[dict[str, Any]], None] | None = None,
) -> tuple[str, ...]:
    """Re-freeze mode drift on a published snapshot after the pre-dispatch deadline (2.6.8).

    Inside the SIGALRM deadline the runners verify with ``lenient_modes`` only. This runs once the
    deadline is disarmed: a cheap lstat walk, then, only if some entry carries a write bit, the
    full verify_snapshot(refreeze=True) -- content and inventory again before any chmod, the
    bounded re-freeze, and ``on_refreeze`` (the journal note) from the process that ran chmod.
    Returns the re-frozen snapshot-relative paths.
    """
    try:
        drifted = snapshot_has_mode_drift(snapshot)
    except SnapshotIntegrityError:
        drifted = True  # a symlink or special file: the full check below reports it, with remedy
    if not drifted:
        return ()
    return _verify_published(snapshot, manifest, None, True, on_refreeze)


def _verify_published(
    target: pathlib.Path, manifest: dict[str, Any], deadline: MonotonicDeadline | None,
    refreeze: bool, on_refreeze: Callable[[dict[str, Any]], None] | None,
    lenient_modes: bool = False,
) -> tuple[str, ...]:
    """verify_snapshot on the published snapshot, with the remedy an integrity fault needs."""
    try:
        return verify_snapshot(target, manifest, deadline, refreeze=refreeze,
                               on_refreeze=on_refreeze, lenient_modes=lenient_modes)
    except SnapshotModeUnstableError:
        raise  # its own remedy: deleting the snapshot cannot help while modes keep changing
    except SnapshotIntegrityError as exc:
        relative = f"{target.parent.parent.name}/snapshots/{target.name}"
        raise SnapshotIntegrityError(
            f"{exc}; remedy: in the audit directory run `{snapshot_delete_remedy(relative)}` (it "
            "is content-addressed and is rebuilt identically from the deposit), then re-run"
        ) from exc


def _snapshot_member(base: pathlib.Path, relative: str, label: str) -> pathlib.Path:
    """beneath() + regular_file() for a manifest row, as integrity faults naming only the row."""
    resolved = (base / relative).resolve()
    try:
        resolved.relative_to(base.resolve())
    except ValueError:
        raise SnapshotIntegrityError(f"{label} escapes the snapshot") from None
    if resolved.is_symlink() or not resolved.is_file():
        raise SnapshotIntegrityError(f"{label} is missing or is not a regular file")
    return resolved


def _snapshot_mode_entries(snapshot: pathlib.Path) -> list[pathlib.Path]:
    """Every entry of the snapshot tree, root included, that carries a write bit.

    A symlink or special file anywhere is an integrity fault: freeze_tree's chmod would follow a
    link, and nothing the snapshot builder writes is either.
    """
    writable: list[pathlib.Path] = []
    for directory, subdirs, files in os.walk(snapshot):
        base = pathlib.Path(directory)
        if stat.S_IMODE(os.lstat(base).st_mode) & 0o222:
            writable.append(base)
        for name in subdirs:
            if os.path.islink(base / name):
                raise SnapshotIntegrityError("the snapshot tree contains a symlink")
        for name in files:
            info = os.lstat(base / name)
            if not stat.S_ISREG(info.st_mode):
                raise SnapshotIntegrityError("the snapshot tree contains a symlink or special file")
            if stat.S_IMODE(info.st_mode) & 0o222:
                writable.append(base / name)
    return writable


_SNAPSHOT_ROOT_ENTRIES = ("snapshot_manifest.json", "raw", "evidence", "context")


def _nfc(value: str) -> str:
    # A synced folder may hand back a name in another Unicode normalization; the filesystem
    # treats both spellings as one file, so the inventory must too.
    return unicodedata.normalize("NFC", value)


def _check_snapshot_inventory(snapshot: pathlib.Path, manifest: dict[str, Any]) -> None:
    """The snapshot holds exactly what its manifest names, and nothing else (2.6.8).

    The per-row checks prove every listed file; this proves no file was added. raw/ must hold
    exactly the source_files paths (it is the provider sandbox's source), evidence/ the artifacts
    plus evidence_index.json, context/ the context rows, and the root only snapshot_manifest.json
    beside those three directories. A directory must lead to a listed file. It runs before any
    chmod, so a file that lands in a drifted directory is refused, never sealed by a re-freeze.
    Messages give counts only: a stray name can be a deposit or participant name.
    """
    expected = {
        "raw": {_nfc(row["path"]) for row in manifest.get("source_files", [])},
        "evidence": {_nfc(row["path"]) for row in
                     manifest.get("common_evidence", {}).get("artifacts", [])}
        | {"evidence_index.json"},
        "context": {_nfc(row["path"]) for row in manifest.get("context", [])},
    }
    try:
        root_names = os.listdir(snapshot)
    except OSError as exc:
        raise SnapshotIntegrityError(f"the snapshot root cannot be listed ({exc.strerror})") from None
    unexpected_root = [name for name in root_names if _nfc(name) not in _SNAPSHOT_ROOT_ENTRIES]
    if unexpected_root:
        raise SnapshotIntegrityError(
            f"immutable snapshot inventory differs: the snapshot root holds {len(unexpected_root)} "
            "entry(ies) besides snapshot_manifest.json, raw, evidence and context"
        )
    manifest_file = snapshot / "snapshot_manifest.json"
    if manifest_file.is_symlink() or not manifest_file.is_file():
        raise SnapshotIntegrityError("snapshot manifest is missing or is not a regular file")
    for area, names in expected.items():
        base = snapshot / area
        if base.is_symlink() or (base.exists() and not base.is_dir()):
            raise SnapshotIntegrityError(f"immutable snapshot inventory differs: {area}/ is not a directory")
        if not base.exists():
            if names:
                raise SnapshotIntegrityError(
                    f"immutable snapshot inventory differs: {area}/ is missing "
                    f"({len(names)} listed file(s))")
            continue
        allowed_dirs = {"/".join(parts[:depth]) for parts in (name.split("/") for name in names)
                        for depth in range(1, len(parts))}
        found: set[str] = set()
        stray_dirs = special = 0

        def unlistable(exc: OSError) -> None:
            raise SnapshotIntegrityError(
                f"immutable snapshot inventory cannot be listed under {area}/ ({exc.strerror})")

        for directory, subdirs, files in os.walk(base, onerror=unlistable):
            here = pathlib.Path(directory)
            for name in subdirs:
                relative = _nfc((here / name).relative_to(base).as_posix())
                if os.path.islink(here / name):
                    special += 1
                elif relative not in allowed_dirs:
                    stray_dirs += 1
            for name in files:
                info = os.lstat(here / name)
                if not stat.S_ISREG(info.st_mode):
                    special += 1
                    continue
                found.add(_nfc((here / name).relative_to(base).as_posix()))
        unlisted, missing = len(found - names), len(names - found)
        if unlisted or missing or stray_dirs or special:
            parts = []
            if unlisted:
                parts.append(f"{unlisted} file(s) the manifest does not list")
            if missing:
                parts.append(f"{missing} listed file(s) missing")
            if stray_dirs:
                parts.append(f"{stray_dirs} directory(ies) that hold no listed file")
            if special:
                parts.append(f"{special} symlink(s) or special file(s)")
            raise SnapshotIntegrityError(
                f"immutable snapshot inventory differs: {area}/ holds {', '.join(parts)}"
            )


def _depth_limited_paths(snapshot: pathlib.Path, paths: list[pathlib.Path]) -> list[str]:
    """Directories at most two deep (raw/pkg, evidence, context): never a deposit name."""
    names = set()
    for path in paths:
        parts = path.relative_to(snapshot).parts
        if not path.is_dir():
            parts = parts[:-1]
        names.add("/".join(parts[:2]) or ".")
    return sorted(names)


_REFREEZE_THREAD_LOCK = threading.Lock()


def _refreeze_snapshot(
    snapshot: pathlib.Path, manifest: dict[str, Any],
    on_refreeze: Callable[[dict[str, Any]], None] | None,
) -> tuple[str, ...]:
    writable = _snapshot_mode_entries(snapshot)
    if not writable:
        return ()
    # The five fan-out roles run in threads and may all see the same drift: one lock per process
    # and an flock on a regular file across processes, so exactly one of them re-freezes and the
    # callback (a journal note) runs once.
    lock_path = snapshot.parent.parent / ".snapshot_refreeze.lock"
    with _REFREEZE_THREAD_LOCK, lock_path.open("a+") as handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            writable = _snapshot_mode_entries(snapshot)
            if not writable:
                return ()
            # The inventory once more under the lock, right before the first chmod: a file that
            # landed in a drifted directory while this process waited must be refused, not sealed.
            _check_snapshot_inventory(snapshot, manifest)
            # No cooperative deadline here, and the callers never run this under the SIGALRM
            # pre-dispatch deadline either (they verify with lenient_modes inside it and re-freeze
            # after it is disarmed): an interrupted walk would leave the snapshot half-frozen with
            # no note for the chmods that did happen.
            freeze_tree(snapshot, None)
            remaining = _snapshot_mode_entries(snapshot)
            passes = 1
            while remaining and passes < SNAPSHOT_REFREEZE_ATTEMPTS:
                # One more File Provider reset is repaired quietly; a mode that keeps changing is
                # reported below as what it is, not as a content change.
                time.sleep(SNAPSHOT_REFREEZE_RETRY_SECONDS * passes)
                _check_snapshot_inventory(snapshot, manifest)
                freeze_tree(snapshot, None)
                remaining = _snapshot_mode_entries(snapshot)
                passes += 1
            if remaining:
                raise SnapshotModeUnstableError(
                    f"{SNAPSHOT_MODE_UNSTABLE}: the Phase-1 snapshot's content verifies against its "
                    f"manifest, but {len(remaining)} path(s) regained a write bit after {passes} "
                    "re-freezes: something on this host keeps changing its modes. Stop the sync "
                    "client, or move the audit off the synced folder, then re-run; deleting the "
                    "snapshot does not help while the modes keep changing"
                )
            # Content once more, strictly, now that nothing is writable: the chmod changed no byte.
            verify_snapshot(snapshot, manifest, None)
            facts = {
                "path_count": len(writable),
                "paths": _depth_limited_paths(snapshot, writable),
            }
            if on_refreeze is not None:
                on_refreeze(facts)
            return tuple(path.relative_to(snapshot).as_posix() for path in writable)
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def verify_snapshot(
    snapshot: pathlib.Path, manifest: dict[str, Any] | None = None,
    deadline: MonotonicDeadline | None = None,
    *,
    refreeze: bool = False,
    on_refreeze: Callable[[dict[str, Any]], None] | None = None,
    lenient_modes: bool = False,
) -> tuple[str, ...]:
    """Verify the frozen snapshot; content and inventory first, then modes.

    Content, inventory, manifest and structure failures raise SnapshotIntegrityError, always
    before any mode is touched. Then modes, one of three ways:

    * strict (the default): only a writable raw/ source file is an error, exactly as before
      2.6.8, so old snapshots are judged as they always were. Nothing on disk changes.
    * ``lenient_modes``: modes are not judged at all. Returns the snapshot-relative paths that
      carry a write bit, i.e. what a re-freeze would change, and changes nothing. For a dry run,
      and inside the SIGALRM pre-dispatch deadline, where a chmod walk must never be interrupted.
    * ``refreeze``: every entry's mode is considered; mode-only drift (the File Provider giving a
      frozen directory its write bit back) is re-frozen in place under a lock, retried a bounded
      number of times, the content re-verified, and ``on_refreeze`` called once with
      depth-limited facts. Returns the re-frozen snapshot-relative paths. A mode that will not
      stay frozen raises SnapshotModeUnstableError.
    """
    if refreeze and lenient_modes:
        raise ValueError("verify_snapshot: refreeze and lenient_modes are exclusive")
    check_deadline(deadline, "verifying source snapshot")
    manifest = manifest or strict_json(regular_file(snapshot / "snapshot_manifest.json", "snapshot manifest"))
    expected_hash = manifest.get("snapshot_sha256")
    body = dict(manifest)
    body.pop("snapshot_sha256", None)
    if expected_hash != sha256_bytes(canonical_bytes(body)):
        raise SnapshotIntegrityError("snapshot manifest hash is invalid")
    writable_sources: list[int] = []
    for index, row in enumerate(manifest.get("source_files", [])):
        check_deadline(deadline, "verifying source snapshot")
        label = f"immutable snapshot source_files[{index}]"
        path = _snapshot_member(snapshot / "raw", row["path"], label)
        if path.stat().st_size != row["bytes"] or sha256_file(path, deadline) != row["sha256"]:
            raise SnapshotIntegrityError(f"immutable snapshot source changed: source_files[{index}]")
        if path.stat().st_mode & stat.S_IWUSR:
            writable_sources.append(index)
    for index, row in enumerate(manifest.get("common_evidence", {}).get("artifacts", [])):
        check_deadline(deadline, "verifying common evidence")
        label = f"common evidence artifacts[{index}]"
        path = _snapshot_member(snapshot / "evidence", row["path"], label)
        if path.stat().st_size != row["bytes"] or sha256_file(path, deadline) != row["sha256"]:
            raise SnapshotIntegrityError(f"common evidence changed: artifacts[{index}]")
    index_path = snapshot / "evidence/evidence_index.json"
    if index_path.is_symlink() or not index_path.is_file():
        raise SnapshotIntegrityError("common evidence index is missing or is not a regular file")
    if sha256_file(index_path, deadline) != manifest.get("common_evidence_index_sha256"):
        raise SnapshotIntegrityError("common evidence index changed")
    for row in manifest.get("context", []):
        check_deadline(deadline, "verifying snapshot context")
        path = snapshot / "context" / row["path"]
        if path.is_symlink() or not path.is_file():
            raise SnapshotIntegrityError(f"snapshot context is missing or irregular: {row['path']}")
        if sha256_file(path, deadline) != row["sha256"]:
            raise SnapshotIntegrityError(f"snapshot context changed: {row['path']}")
    check_deadline(deadline, "verifying the snapshot inventory")
    _check_snapshot_inventory(snapshot, manifest)
    if refreeze:
        return _refreeze_snapshot(snapshot, manifest, on_refreeze)
    if lenient_modes:
        return tuple(path.relative_to(snapshot).as_posix() or "."
                     for path in _snapshot_mode_entries(snapshot))
    if writable_sources:
        raise Phase1Error(
            f"immutable snapshot source is writable: source_files[{writable_sources[0]}] "
            "(content unchanged; a synced folder re-applying a mode is the usual cause -- "
            "the runners re-freeze it themselves, or chmod -R a-w the snapshot)"
        )
    return ()


def snapshot_has_mode_drift(snapshot: pathlib.Path) -> bool:
    """A cheap walk (lstat only, no hashing): does any snapshot entry carry a write bit?

    The runners use it after the pre-dispatch deadline to decide whether the full
    verify_snapshot(refreeze=True) is needed at all.
    """
    return bool(_snapshot_mode_entries(snapshot))


def source_hashes(manifest: dict[str, Any]) -> dict[str, str]:
    return {row["path"]: row["sha256"] for row in manifest["source_files"]}


def make_role_packet(role: str, manifest: dict[str, Any], policy: dict[str, Any]) -> dict[str, Any]:
    spec = ROLE_SPECS[role]
    evidence_by_name = {row["path"]: row for row in manifest["common_evidence"]["artifacts"]}
    selected = [evidence_by_name[name] for name in spec["evidence"]]
    return {
        "schema_version": "i4r-phase1-role-packet-1",
        "role": role,
        "area": spec["area"],
        "area_name": spec["name"],
        "source_snapshot_sha256": manifest["snapshot_sha256"],
        "source_manifest": manifest["source_files"],
        "assigned_evidence": selected,
        "raw_access": {"paper": "raw/paper", "package": "raw/pkg", "read_only": True},
        "review_paths": {
            "conclusion_changing": ["specialist", "completeness"],
            "disclosure_critical": ["fairness", "completeness"],
        },
        "model": {
            "requested_model": policy["requested_model"],
            "spawn_alias": policy["spawn_alias"],
            "reasoning_effort": policy["reasoning_effort"],
        },
        "output_contract": {
            "report": "REPORT.md",
            "result": "result.json",
            "evidence_manifest": "evidence/manifest.json",
            "canonical_report": spec["report"],
        },
    }


# The CLI compiles sandbox.filesystem entries as glob-ish patterns, so a rule whose literal
# path contains one of these characters compiles to a pattern that matches nothing.  Every
# allowRead/allowWrite rule silently grants nothing and seatbelt denies every Bash read and
# write under the audit root; the Write tool separately refuses such a path outright with
# "Glob patterns are not allowed in write operations".  The denyRead/denyWrite rules go inert
# at the same moment, so the stated protection of paper/, pkg/ and run.json stops existing too.
# Measured against the real CLI: "?", "*", "[", "]", "{", "}" and "\" all break; space, ":",
# "~", "(", ")" and "!" are safe.  A paper folder named "Can automatic reactions mirror
# exercise dependence?" spent 5,443,051 tokens and $6.43 across five roles and produced no
# report at all, because nothing in the pipeline could see that the sandbox was inert.
SANDBOX_RULE_METACHARACTERS = "?*[]{}\\"


def sandbox_rule_metacharacters(path: pathlib.Path | str) -> str:
    """Return the glob metacharacters in ``path``, in the order they appear."""
    text = str(path)
    return "".join(sorted({character for character in text if character in SANDBOX_RULE_METACHARACTERS}))


def assert_sandbox_expressible(paths: Iterable[pathlib.Path | str], audit_root: pathlib.Path) -> None:
    """Refuse paths that cannot be expressed as a sandbox filesystem rule.

    This runs before any provider call.  The alternative is a run in which every sandboxed
    Bash command is denied while the preflight receipt and the capability probe both report
    success, which is exactly how the failure above went undetected until the budget was gone.
    """
    for path in paths:
        found = sandbox_rule_metacharacters(path)
        if not found:
            continue
        offenders = " ".join(repr(character) for character in found)
        raise Phase1Error(
            f"audit path cannot be expressed as a sandbox rule: {path}\n"
            f"  offending character(s): {offenders}\n"
            f"  The Claude CLI reads sandbox.filesystem entries as glob patterns, so this rule\n"
            f"  matches nothing and every sandboxed Bash read and write is denied.\n"
            f"  Forbidden: {' '.join(SANDBOX_RULE_METACHARACTERS)}   Safe: spaces : ~ ( ) !\n"
            f"  Remedy: the paper and package are copied into the audit dir, so it does not have\n"
            f"  to sit inside the paper folder. Re-run setup_workspace.sh with a clean audit dir\n"
            f"  and I4R_PAPER_DIR set to the real paper folder, which keeps the reservation under\n"
            f"  the paper's true name.\n"
            f"  Audit root: {audit_root}"
        )


ExtraDenyRead = Union[Iterable[pathlib.Path], Callable[["tuple[pathlib.Path, ...]"], Iterable[pathlib.Path]]]


def resolve_extra_deny_read(
    extra_deny_read: ExtraDenyRead, read_roots: Iterable[pathlib.Path],
) -> tuple[pathlib.Path, ...]:
    """The extra read denials for one attempt.

    A fixed iterable is used as given. A callable is asked with the attempt's own read roots
    (allowRead: source, packet, output), so a denial list can depend on which attempt it is --
    the 2.6.4 verifier denies every sibling attempt under phase1/roles except its own.
    """
    if callable(extra_deny_read):
        return tuple(pathlib.Path(path) for path in extra_deny_read(tuple(read_roots)))
    return tuple(pathlib.Path(path) for path in extra_deny_read)


def sandbox_settings(
    audit_root: pathlib.Path,
    phase_root: pathlib.Path,
    snapshot: pathlib.Path,
    role_packet_dir: pathlib.Path,
    output: pathlib.Path,
    *,
    extra_deny_read: ExtraDenyRead = (),
) -> dict[str, Any]:
    """The exact settings one attempt runs under.

    ``extra_deny_read`` appends read denials after the protected list (the 2.6.4 verifier uses it
    to keep the coordinator's reasoning, the settled/critic artefacts and every other attempt's
    output out of its view). It may be a callable of the attempt's read roots
    (resolve_extra_deny_read). The default adds nothing, so an auditor's settings, and every hash
    of them, are unchanged.
    """
    extra = list(resolve_extra_deny_read(extra_deny_read, (snapshot, role_packet_dir, output)))
    assert_sandbox_expressible([audit_root, snapshot, role_packet_dir, output, *extra], audit_root)
    protected = [
        audit_root / "paper", audit_root / "pkg", audit_root / "report", audit_root / "work",
        audit_root / "critique", audit_root / "run.json", audit_root / "run.events.jsonl",
        audit_root / "invocations.jsonl", audit_root / "package_manifest.json",
        audit_root / "pipeline-policy.snapshot.json",
    ]
    return {
        "permissions": {
            "allow": [*NON_BASH_TOOLS, *BASH_ALLOWLIST],
            "deny": ["WebFetch", "WebSearch"],
        },
        "sandbox": {
            "enabled": True,
            "failIfUnavailable": True,
            "autoAllowBashIfSandboxed": False,
            "allowUnsandboxedCommands": False,
            "network": {"allowedDomains": [], "deniedDomains": ["*"]},
            "filesystem": {
                "allowRead": [str(snapshot), str(role_packet_dir), str(output)],
                "allowWrite": [str(output)],
                "denyRead": [str(path) for path in [*protected, *extra]],
                "denyWrite": [str(path) for path in [*protected, snapshot, role_packet_dir]],
            },
        },
        "i4rBashAllowlist": list(BASH_ALLOWLIST),
    }


def sandbox_policy_hash() -> str:
    abstract = {
        "safe_mode": True,
        # Bare mode refuses OAuth/keychain credentials, so a subscription-authenticated
        # CLI cannot run under it; run_phase1.claude_command does not pass --bare and this
        # hash must say so, or the bound contract asserts a mode the CLI is not in.
        "bare": False,
        "network": {"allowedDomains": [], "deniedDomains": ["*"]},
        "tools": list(NON_BASH_TOOLS) + ["Bash"],
        "bash_allowlist": list(BASH_ALLOWLIST),
        "no_session_persistence": True,
        "plugins": False,
        "mcp_servers": {},
    }
    return sha256_bytes(canonical_bytes(abstract))


def validate_sandbox_contract(
    settings: dict[str, Any], snapshot: pathlib.Path, role_packet_dir: pathlib.Path,
    output: pathlib.Path, audit_root: pathlib.Path,
) -> None:
    sandbox = settings.get("sandbox")
    if not isinstance(sandbox, dict) or sandbox.get("enabled") is not True or sandbox.get("failIfUnavailable") is not True:
        raise Phase1Error("sandbox must be enabled and fail if unavailable")
    if sandbox.get("autoAllowBashIfSandboxed") is not False or sandbox.get("allowUnsandboxedCommands") is not False:
        raise Phase1Error("sandbox permits unreviewed or unsandboxed Bash")
    if sandbox.get("network") != {"allowedDomains": [], "deniedDomains": ["*"]}:
        raise Phase1Error("sandbox network policy is not deny-all")
    filesystem = sandbox.get("filesystem") or {}
    # Checked here as well as at emission: this validator is the contract's last word before a
    # provider call, and comparing the emitted rules against themselves cannot see a rule that
    # is well-formed JSON but matches nothing.
    assert_sandbox_expressible(
        [
            *filesystem.get("allowRead", []),
            *filesystem.get("allowWrite", []),
            *filesystem.get("denyRead", []),
            *filesystem.get("denyWrite", []),
        ],
        audit_root,
    )
    expected_read = [str(snapshot), str(role_packet_dir), str(output)]
    if filesystem.get("allowRead") != expected_read or filesystem.get("allowWrite") != [str(output)]:
        raise Phase1Error("sandbox read/write paths differ from the exact Phase-1 contract")
    denied_read = {pathlib.Path(value).resolve() for value in filesystem.get("denyRead", [])}
    if audit_root.resolve() / "pkg" not in denied_read or audit_root.resolve() / "paper" not in denied_read:
        raise Phase1Error("sandbox does not deny live raw sources")
    permissions = settings.get("permissions") or {}
    if permissions.get("allow") != [*NON_BASH_TOOLS, *BASH_ALLOWLIST]:
        raise Phase1Error("Bash/tool allowlist differs from the hashed Phase-1 contract")
    if settings.get("i4rBashAllowlist") != list(BASH_ALLOWLIST):
        raise Phase1Error("settings do not bind the exact Bash allowlist")


def provider_free_preflight(
    audit_root: pathlib.Path,
    phase_root: pathlib.Path,
    snapshot: pathlib.Path,
    manifest: dict[str, Any],
    role_packets: dict[str, pathlib.Path],
    deadline: MonotonicDeadline | None = None,
    *,
    destinations: Iterable[tuple[pathlib.Path, str]] | None = None,
    workspace: pathlib.Path | None = None,
    extra_deny_read: ExtraDenyRead = (),
    refreeze: bool = False,
    on_refreeze: Callable[[dict[str, Any]], None] | None = None,
    lenient_modes: bool = False,
) -> dict[str, Any]:
    """Exercise exact source, sandbox, output, and promotion paths without a provider process.

    The keyword arguments exist for the 2.6.4 verifier, which must preflight its own promotion
    target in its own workspace (running this as-is would overwrite the Phase-1 receipt) under
    its own extra read denials (a fixed iterable or a callable of the read roots, exactly as
    sandbox_settings takes them). Their defaults are the five auditors' destinations, the Phase-1
    workspace ``phase1/preflight/<snapshot>`` and no extra denial, so an auditor preflight and its
    receipt are unchanged. ``refreeze``, ``on_refreeze`` and ``lenient_modes`` go to
    verify_snapshot. This runs inside the SIGALRM pre-dispatch deadline, so the runners pass
    ``lenient_modes`` and re-freeze only after the deadline is disarmed.
    """
    if not callable(extra_deny_read):
        extra_deny_read = tuple(extra_deny_read)
    # First, before any other work: a root that cannot be expressed as a sandbox rule makes
    # every later check meaningless, because the checks below write through Python rather than
    # through the sandbox and therefore succeed no matter what the sandbox would have done.
    assert_sandbox_expressible([audit_root, phase_root, snapshot], audit_root)
    verify_snapshot(snapshot, manifest, deadline, refreeze=refreeze, on_refreeze=on_refreeze,
                    lenient_modes=lenient_modes)
    validate_destination(audit_root, phase_root, "Phase-1 runtime directory")
    if destinations is None:
        for role, spec in ROLE_SPECS.items():
            check_deadline(deadline, "checking Phase-1 destinations")
            validate_destination(audit_root, audit_root / spec["report"], f"{role} report destination")
            validate_destination(audit_root, audit_root / spec["work"], f"{role} evidence destination")
    else:
        for destination, label in destinations:
            check_deadline(deadline, "checking Phase-1 destinations")
            validate_destination(audit_root, destination, label)
    smoke = workspace if workspace is not None else phase_root / "preflight" / manifest["snapshot_sha256"]
    validate_destination(audit_root, smoke, "preflight workspace")
    if smoke.exists():
        if smoke.is_symlink() or not smoke.is_dir():
            raise Phase1Error(f"unsafe preflight path: {smoke}")
    else:
        smoke.mkdir(parents=True)
    contract_hashes: dict[str, str] = {}
    for role, packet_dir in role_packets.items():
        check_deadline(deadline, f"preflighting {role}")
        validate_destination(audit_root, packet_dir, f"{role} role packet")
        output = smoke / role / "output.partial"
        validate_destination(audit_root, output, f"{role} preflight output")
        output.mkdir(parents=True, exist_ok=True)
        provider_source = snapshot / "raw"
        settings = sandbox_settings(
            audit_root, phase_root, provider_source, packet_dir, output,
            extra_deny_read=extra_deny_read,
        )
        validate_sandbox_contract(settings, provider_source, packet_dir, output, audit_root)
        # Read the exact provider-visible roots.
        for source_name in ("paper", "pkg"):
            source_dir = provider_source / source_name
            if source_dir.is_symlink() or not source_dir.is_dir():
                raise Phase1Error(f"preflight {source_name} source must be a non-symlink directory")
        regular_file(packet_dir / "ROLE_PACKET.json", "preflight role packet").read_bytes()
        # Exercise the exact output and atomic-promotion filesystems.  The smoke target is private;
        # canonical reports are not touched by preflight.
        partial = output / "smoke.partial"
        partial.write_text("phase1-preflight\n", encoding="utf-8")
        with partial.open("rb") as handle:
            os.fsync(handle.fileno())
        promoted = smoke / role / "smoke.promoted"
        os.replace(partial, promoted)
        if promoted.read_text(encoding="utf-8") != "phase1-preflight\n":
            raise Phase1Error("preflight output promotion changed bytes")
        promoted.unlink()
        contract_hashes[role] = sha256_bytes(canonical_bytes(settings))
    receipt = {
        "schema_version": "i4r-phase1-preflight-1",
        "status": "passed",
        "provider_calls": 0,
        "snapshot_sha256": manifest["snapshot_sha256"],
        "sandbox_policy_sha256": sandbox_policy_hash(),
        "role_sandbox_settings_sha256": contract_hashes,
        "checks": {
            "source_accessibility": True,
            "source_hashes": True,
            "sandbox_paths": True,
            # The checks in this receipt write through Python, not through the sandbox, so
            # they cannot tell a working sandbox from an inert one.  This is the one property
            # of the sandbox rules that is decidable without a provider call.
            "sandbox_rules_expressible": True,
            "bash_allowlist": True,
            "network_denied": True,
            "output_write": True,
            "atomic_promotion": True,
        },
    }
    atomic_write(smoke / "receipt.json", pretty_bytes(receipt))
    return receipt


def _validated_scientific_fields(
    scientific: Any, snapshot_manifest: dict[str, Any],
) -> dict[str, Any]:
    """Validate the complete scientific record that a repair must never change."""
    if not isinstance(scientific, dict):
        raise OutputValidationError("scientific_fields must be an object")
    if set(scientific) != {"passed_check_ids", "findings", "open_question_ids"}:
        raise OutputValidationError("scientific_fields has an invalid exact field set")
    for name in ("passed_check_ids", "open_question_ids"):
        values = scientific[name]
        if (
            not isinstance(values, list)
            or any(not isinstance(value, str) or not value for value in values)
            or len(values) != len(set(values))
        ):
            raise OutputValidationError(f"scientific_fields.{name} must contain unique IDs")
    findings = scientific["findings"]
    if not isinstance(findings, list):
        raise OutputValidationError("scientific_fields.findings must be an array")
    known = source_hashes(snapshot_manifest)
    finding_ids: set[str] = set()
    required = {
        "id", "status", "claim", "evidence", "impact", "remedy",
        "conclusion_changing", "disclosure_critical",
    }
    for index, finding in enumerate(findings):
        if not isinstance(finding, dict) or set(finding) != required:
            raise OutputValidationError(f"scientific_fields.findings[{index}] is invalid")
        if (
            not isinstance(finding["id"], str) or not finding["id"]
            or finding["id"] in finding_ids
            or finding["status"] not in {"CONFIRMED", "LIKELY", "SPECULATIVE"}
            or any(
                not isinstance(finding[name], str) or not finding[name].strip()
                for name in ("claim", "impact", "remedy")
            )
            or not isinstance(finding["conclusion_changing"], bool)
            or not isinstance(finding["disclosure_critical"], bool)
        ):
            raise OutputValidationError(f"scientific_fields.findings[{index}] has invalid values")
        evidence = finding["evidence"]
        if not isinstance(evidence, list) or not evidence:
            raise OutputValidationError(
                f"scientific_fields.findings[{index}].evidence must be non-empty"
            )
        for evidence_index, item in enumerate(evidence):
            if not isinstance(item, dict) or set(item) != {
                "source_path", "source_sha256", "result",
            }:
                raise OutputValidationError(
                    f"scientific_fields.findings[{index}].evidence[{evidence_index}] is invalid"
                )
            if (
                known.get(item.get("source_path")) != item.get("source_sha256")
                or not isinstance(item.get("result"), str)
                or not item["result"].strip()
            ):
                raise OutputValidationError(
                    f"scientific_fields.findings[{index}].evidence[{evidence_index}] "
                    "is not bound to a source/result"
                )
        finding_ids.add(finding["id"])
    return scientific


def frozen_evidence_artifacts(evidence_root: pathlib.Path) -> list[dict[str, Any]]:
    """Every evidence artifact (not the manifest) with its size and hash, for a repair freeze.

    Shared by the auditor projection below and the 2.6.4 verifier projection, so both freeze the
    evidence tree the same way.
    """
    artifacts: list[dict[str, Any]] = []
    for path in sorted(evidence_root.rglob("*"), key=lambda item: item.as_posix()):
        if path.is_symlink() or (not path.is_file() and not path.is_dir()):
            raise OutputValidationError("role evidence contains an unsafe artifact", rule="output_symlink")
        if path.is_file() and path.name != "manifest.json":
            artifacts.append({
                "path": path.relative_to(evidence_root).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            })
    return artifacts


def scientific_projection(
    output: pathlib.Path, role: str, snapshot_manifest: dict[str, Any],
) -> dict[str, Any]:
    """Capture the byte-frozen scientific surface before a focused output repair.

    The evidence manifest itself may be corrected by a format-only repair, but every evidence
    artifact it describes, the complete report bytes, and every structured finding field are
    immutable.  This prevents a nominal formatting repair from becoming an unreviewed re-audit.
    """
    report = regular_file(output / "REPORT.md", "role report")
    result = strict_json(regular_file(output / "result.json", "role result"))
    if (
        result.get("schema_version") != ROLE_RESULT_VERSION
        or result.get("role") != role
        or result.get("status") != "complete"
        or result.get("completion_marker") is not True
        or result.get("source_snapshot_sha256") != snapshot_manifest.get("snapshot_sha256")
    ):
        raise OutputValidationError("role result is not eligible for focused repair")
    scientific = _validated_scientific_fields(
        result.get("scientific_fields"), snapshot_manifest,
    )
    evidence_root = output / "evidence"
    regular_file(evidence_root / "manifest.json", "role evidence manifest")
    artifacts = frozen_evidence_artifacts(evidence_root)
    if not artifacts:
        raise OutputValidationError("role result has no evidence artifacts to freeze")
    finding_hashes = {
        finding["id"]: {
            name: sha256_bytes(canonical_bytes(finding[name]))
            for name in ("claim", "evidence", "impact", "remedy")
        }
        for finding in scientific["findings"]
    }
    report_bytes = report.read_bytes()
    return {
        "schema_version": "i4r-phase1-scientific-projection-1",
        "role": role,
        "source_snapshot_sha256": snapshot_manifest["snapshot_sha256"],
        "scientific_fields": scientific,
        "finding_field_sha256": finding_hashes,
        "report_prose": {
            "path": "REPORT.md", "bytes": len(report_bytes),
            "sha256": sha256_bytes(report_bytes),
        },
        "evidence_artifacts": artifacts,
    }


FINDING_ELEMENTS = ("evidence", "benign", "impact", "remedy")


def findings_elements_missing(finding_text: str) -> list[str]:
    """The four required elements a Findings section does not name.

    Extracted so the rule can be tested directly. On the sixth production audit one role lost two
    of its three attempts to this check: its reports carried *Benign explanations tested*, *Impact*
    and *Remedy* labels and substantive evidence, but never the word "evidence" inside the
    Findings section, and the message said only that the four elements were "omitted". The prompt
    now spells the four labels, and this message names the one that is missing.
    """
    lower = finding_text.lower()
    return [term for term in FINDING_ELEMENTS if term not in lower]


@dataclass(frozen=True)
class OutputContract:
    """How one role class's output is admitted and frozen for a focused repair.

    ``validate(output, role, snapshot_manifest, *, expected_scientific_projection=None)`` returns
    a ValidatedRoleOutput or raises OutputValidationError; ``projection(output, role,
    snapshot_manifest)`` returns the byte-frozen surface a repair must preserve. The auditor
    contract below is exactly validate_role_output / scientific_projection. The 2.6.4 verifier
    builds its own contract (scripts/run_verifier.py) from the generic admission checks that
    follow, so both role classes share one implementation of manifest closure, source joins,
    the leak, participant and absolute-path rules and the closed file set.
    """

    name: str
    validate: Callable[..., "ValidatedRoleOutput"]
    projection: Callable[..., dict[str, Any]]


# ------------------------------------------------------------------ generic admission checks
# Each helper raises exactly the message the auditor validator has always raised, in the same
# order, and tags it with the rule it enforces.

def read_role_output(
    output: pathlib.Path,
) -> tuple[pathlib.Path, pathlib.Path, pathlib.Path, str, dict[str, Any]]:
    """The three contract files, the report text and the parsed result."""
    with _admission_rule("output_files"):
        report = regular_file(output / "REPORT.md", "role report")
        result_path = regular_file(output / "result.json", "role result")
        evidence_manifest_path = regular_file(output / "evidence/manifest.json", "role evidence manifest")
    try:
        report_text = report.read_text(encoding="utf-8")
    except UnicodeError as exc:
        raise OutputValidationError("role report is not UTF-8", rule="report_utf8") from exc
    with _admission_rule("result_json"):
        result = strict_json(result_path)
    return report, result_path, evidence_manifest_path, report_text, result


def check_result_bindings(
    result: dict[str, Any], report: pathlib.Path, evidence_manifest_path: pathlib.Path,
) -> None:
    """result.json binds the final REPORT.md and evidence-manifest bytes (repairable)."""
    if result.get("report") != {"path": "REPORT.md", "sha256": sha256_file(report)}:
        raise OutputValidationError(
            "role result report hash is invalid", repairable=True, rule="report_hash")
    if result.get("evidence_manifest") != {
        "path": "evidence/manifest.json", "sha256": sha256_file(evidence_manifest_path)
    }:
        raise OutputValidationError(
            "role result evidence-manifest hash is invalid", repairable=True, rule="manifest_hash")


def snapshot_deposit_exemptions(snapshot: pathlib.Path) -> frozenset[str]:
    """Code-shaped variable names of the deposit frozen in a Phase-1 snapshot (raw/pkg)."""
    return provider_safety.deposit_exemptions(snapshot / "raw" / "pkg")


def check_output_hygiene(
    report_text: str, result: dict[str, Any], deposit_exemptions: Iterable[str] = frozenset(),
    *, participant_code_rule: bool = True,
) -> None:
    """No private absolute path or participant code in REPORT.md or result.json; no participant
    id beside a date.

    The participant-code rule is the coordinator's own (provider_safety.py). Before 2.6.5 only the
    coordinator applied it, so a Phase-1 report that quoted a code-shaped token was admitted and
    paid for here and then refused, unrecoverably, as coordinator input. ``deposit_exemptions``
    carries the deposit's own code-shaped column names (snapshot_deposit_exemptions), which both
    stages admit. ``participant_code_rule=False`` is only for re-admitting a verifier record made
    before 2.6.5, whose prompt did not state the rule (run_verifier._readmit)."""
    if ABSOLUTE_PATH.search(report_text):
        # NOT repairable, although it was marked so for a good reason: a cold re-sample throws
        # away a completed review to fix a path string. But the focused repair byte-freezes
        # REPORT.md (scientific_projection.report_prose.sha256) and compels "copy REPORT.md
        # byte-for-byte", so a repair that fixes the path is rejected for changing frozen bytes
        # and a repair that obeys is rejected for the same leak: a guaranteed-fail paid call.
        # Making leaks warm-repairable needs a repair mode that freezes finding identity and
        # structure but permits masking inside prose; until that exists, say so and re-sample.
        raise OutputValidationError(
            "role report leaks a private absolute path (not focused-repairable: the repair "
            "contract freezes REPORT.md byte-for-byte)", rule="report_absolute_path")
    # result.json was never scanned, only REPORT.md. On the second production audit the
    # code-and-methods role kept its report clean -- it described the deposit's hardcoded
    # setwd() paths without quoting them -- but wrote the literal path into the structured
    # finding, which is what feeds the draft registry. The coordinator's provider-safety gate
    # then refused the entire 27-row input, one leak costing every finding in the audit.
    result_leak = ABSOLUTE_PATH.search(json.dumps(result, ensure_ascii=False))
    if result_leak:
        # Same as the report leak above: the leak sits inside scientific_fields, which the
        # focused repair must copy byte-equivalent, so the repair cannot fix it.
        raise OutputValidationError(
            "role result.json leaks a private absolute path: "
            f"{result_leak.group(0)!r}. Mask a path found inside the deposit as "
            "<windows-user-dir>/... or <unix-user-dir>/... and keep the rest "
            "(not focused-repairable: scientific_fields is frozen byte-equivalent)",
            rule="result_absolute_path",
        )
    for number, line in enumerate(report_text.splitlines(), 1):
        if participant_date_linked(line):
            raise OutputValidationError(
                f"role report line {number} links a participant identifier to a date/time",
                rule="participant_date")
    # Not focused-repairable, for the reason given for absolute paths above: the repair contract
    # freezes REPORT.md and scientific_fields byte-for-byte, so only a fresh sample can drop it.
    if not participant_code_rule:
        return
    hit = provider_safety.first_participant_code(report_text, deposit_exemptions)
    if hit:
        raise OutputValidationError(
            f"role report line {hit.line} writes a participant-code-shaped token that is not a "
            "column name in the deposit's data files; describe participants by counts or row "
            "positions (not focused-repairable: the repair contract freezes REPORT.md)",
            rule="report_participant_code")
    if provider_safety.first_participant_code(json.dumps(result, ensure_ascii=False), deposit_exemptions):
        raise OutputValidationError(
            "role result.json writes a participant-code-shaped token that is not a column name in "
            "the deposit's data files (not focused-repairable: scientific_fields is frozen "
            "byte-equivalent)",
            rule="result_participant_code")


def check_report_joins(report_text: str, known: dict[str, str]) -> tuple[tuple[str, str], ...]:
    """At least one <!-- source:path@sha256 --> join, every one bound to the snapshot."""
    bindings = tuple((path.strip(), digest.lower()) for path, digest in REPORT_SOURCE_BINDING.findall(report_text))
    if not bindings:
        raise OutputValidationError("role report has no source-hash join", rule="report_join_missing")
    for path, digest in bindings:
        if known.get(path) != digest:
            # The citation lives in REPORT.md, which the focused repair freezes byte-for-byte.
            raise OutputValidationError(
                f"role report cites an unbound source hash: {path} (not focused-repairable: "
                "the citation is in REPORT.md, which the repair contract freezes)",
                rule="report_join_unbound")
    return bindings


def check_evidence_manifest(
    evidence_manifest_path: pathlib.Path, role: str, snapshot_manifest: dict[str, Any],
    known: dict[str, str], *, schema_version: str = ROLE_RESULT_VERSION,
) -> dict[str, Any]:
    """The manifest header and its non-empty, snapshot-bound source bindings."""
    with _admission_rule("manifest_header"):
        manifest = strict_json(evidence_manifest_path)
    if (
        manifest.get("schema_version") != schema_version
        or manifest.get("role") != role
        or manifest.get("source_snapshot_sha256") != snapshot_manifest["snapshot_sha256"]
        or not isinstance(manifest.get("artifacts"), list)
        or not manifest["artifacts"]
    ):
        raise OutputValidationError("role evidence manifest header is invalid", rule="manifest_header")
    manifest_bindings = manifest.get("source_bindings")
    if not isinstance(manifest_bindings, list) or not manifest_bindings:
        raise OutputValidationError(
            "role evidence manifest has no source bindings", rule="manifest_bindings_missing")
    for binding in manifest_bindings:
        if not isinstance(binding, dict) or known.get(binding.get("path")) != binding.get("sha256"):
            raise OutputValidationError(
                "role evidence manifest contains an unbound source hash", repairable=True,
                rule="manifest_binding_unbound")
    return manifest


def check_output_tree(output: pathlib.Path, manifest: dict[str, Any]) -> set[str]:
    """No symlink or special file; every artifact row exact; closure; the closed file set.

    Returns the evidence-relative artifact paths the manifest lists.
    """
    evidence_root = output / "evidence"
    for path in output.rglob("*"):
        if path.is_symlink():
            raise OutputValidationError(
                f"role output contains a symlink: {path.relative_to(output)}", rule="output_symlink")
        if not path.is_file() and not path.is_dir():
            raise OutputValidationError(
                f"role output contains a special file: {path.relative_to(output)}", rule="output_symlink")
    seen: set[str] = set()
    for item in manifest["artifacts"]:
        if not isinstance(item, dict) or not isinstance(item.get("path"), str):
            raise OutputValidationError(
                "role evidence manifest contains an invalid artifact row", rule="artifact_row")
        relative = item["path"]
        if relative == "manifest.json" or relative in seen:
            raise OutputValidationError(
                f"duplicate or recursive evidence path: {relative}", rule="artifact_row")
        canonical = posixpath.normpath(relative)
        if canonical != relative or relative.startswith(("./", "/")):
            # "./aggregate_check.json" passes beneath(), regular_file() and the hash check, then
            # fails the closure set-comparison below with a message about a MISSING file. Name
            # the actual fault; it is a structural JSON fix, so the focused repair may do it.
            raise OutputValidationError(
                f"role evidence artifact path must be the plain relative path {canonical!r}, "
                f"not {relative!r}", repairable=True, rule="artifact_path")
        seen.add(relative)
        with _admission_rule("artifact_row"):
            path = beneath(evidence_root, evidence_root / relative, "role evidence")
            regular_file(path, "role evidence")
        if path.stat().st_size != item.get("bytes") or sha256_file(path) != item.get("sha256"):
            raise OutputValidationError(
                f"role evidence hash/size mismatch: {relative}", rule="artifact_hash")
    extras = sorted(
        path.relative_to(evidence_root).as_posix()
        for path in evidence_root.rglob("*")
        if path.is_file() and path.name != "manifest.json"
    )
    if extras != sorted(seen):
        raise OutputValidationError(
            "role evidence manifest does not close over every evidence file", rule="manifest_closure")
    top_files = sorted(
        path.relative_to(output).as_posix() for path in output.rglob("*") if path.is_file()
    )
    expected_files = sorted(["REPORT.md", "result.json", "evidence/manifest.json", *[
        f"evidence/{name}" for name in seen
    ]])
    if top_files != expected_files:
        raise OutputValidationError(
            "role output contains files outside the closed output contract", rule="closed_file_set")
    return seen


def validate_role_output(
    output: pathlib.Path,
    role: str,
    snapshot_manifest: dict[str, Any],
    *,
    expected_scientific_projection: dict[str, Any] | None = None,
    deposit_exemptions: Iterable[str] = frozenset(),
) -> ValidatedRoleOutput:
    """Admit one role output, or reject it as an invalid sample.

    Every fault reachable from the provider's own output tree -- a REPORT.md that was never
    written, a result.json that is not JSON, an evidence path that escapes its root -- is a
    property of this sample, not of the runner. The bare Phase1Error those helpers raise fell
    straight through run_role's `except OutputValidationError`, so the attempt was neither
    recorded nor re-sampled: it consumed a full provider call and wrote no ledger row at all.
    """
    try:
        return _validated_role_output(
            output, role, snapshot_manifest,
            expected_scientific_projection=expected_scientific_projection,
            deposit_exemptions=deposit_exemptions,
        )
    except OutputValidationError:
        raise
    except Phase1Error as exc:
        raise OutputValidationError(str(exc)) from exc


def _validated_role_output(
    output: pathlib.Path,
    role: str,
    snapshot_manifest: dict[str, Any],
    *,
    expected_scientific_projection: dict[str, Any] | None = None,
    deposit_exemptions: Iterable[str] = frozenset(),
) -> ValidatedRoleOutput:
    report, result_path, evidence_manifest_path, report_text, result = read_role_output(output)
    scientific = result.get("scientific_fields")
    structurally_usable = (
        result.get("schema_version") == ROLE_RESULT_VERSION
        and result.get("role") == role
        and result.get("status") == "complete"
        and result.get("completion_marker") is True
        and isinstance(scientific, dict)
    )
    if not structurally_usable:
        raise OutputValidationError("role result is not a scientifically usable completed result")
    scientific = _validated_scientific_fields(scientific, snapshot_manifest)
    findings = scientific["findings"]
    if result.get("source_snapshot_sha256") != snapshot_manifest["snapshot_sha256"]:
        raise OutputValidationError("role result does not bind the immutable source snapshot")
    check_result_bindings(result, report, evidence_manifest_path)
    if len(report_text.strip()) < 1000 or PLACEHOLDER.search(report_text):
        raise OutputValidationError("role report is a skeleton or contains placeholders")
    headings = (
        "Result summary", "Brief corrections", "Checks performed and passed", "Findings",
        "Tested hypotheses not confirmed", "Open questions", "Evidence inventory",
    )
    missing = [heading for heading in headings if not re.search(rf"(?im)^##\s+{re.escape(heading)}\s*$", report_text)]
    if missing:
        raise OutputValidationError("role report lacks sections: " + ", ".join(missing))
    marker = f"<!-- i4r-role-complete:{role} -->"
    if report_text.count(marker) != 1:
        raise OutputValidationError("role report lacks its unique completion marker")
    findings_section = re.search(
        r"(?ims)^##\s+Findings\s*$\n(.*?)(?=^##\s+)", report_text
    )
    finding_text = findings_section.group(1) if findings_section else ""
    if findings and not re.search(r"\b(?:CONFIRMED|LIKELY|SPECULATIVE)\b", finding_text):
        raise OutputValidationError("role report findings omit status labels")
    missing_elements = findings_elements_missing(finding_text) if findings else []
    if missing_elements:
        raise OutputValidationError(
            "role report findings omit the labelled element(s): " + ", ".join(missing_elements)
            + " -- each finding must carry *Evidence*:, *Benign explanations tested*:, *Impact*: "
            "and *Remedy*: (the check is for those words inside the Findings section)"
        )
    check_output_hygiene(report_text, result, deposit_exemptions)
    known = source_hashes(snapshot_manifest)
    bindings = check_report_joins(report_text, known)
    manifest = check_evidence_manifest(evidence_manifest_path, role, snapshot_manifest, known)
    check_output_tree(output, manifest)
    if expected_scientific_projection is not None:
        observed_projection = scientific_projection(output, role, snapshot_manifest)
        if canonical_bytes(observed_projection) != canonical_bytes(expected_scientific_projection):
            raise OutputValidationError(
                "focused repair changed the byte-frozen scientific projection",
                repairable=False,
            )
    return ValidatedRoleOutput(
        role=role,
        report=report,
        evidence_manifest=evidence_manifest_path,
        result=result_path,
        report_sha256=sha256_file(report),
        evidence_manifest_sha256=sha256_file(evidence_manifest_path),
        result_sha256=sha256_file(result_path),
        scientific_fields=scientific,
        source_bindings=bindings,
    )


AUDITOR_OUTPUT_CONTRACT = OutputContract(
    name="auditor", validate=validate_role_output, projection=scientific_projection,
)


def role_input_hash(packet: dict[str, Any], prompt: str, probe_key: str) -> str:
    return sha256_bytes(canonical_bytes({
        "packet": packet,
        "prompt_sha256": sha256_bytes(prompt.encode("utf-8")),
        "probe_key": probe_key,
        "sandbox_policy_sha256": sandbox_policy_hash(),
        "runner_version": PHASE1_VERSION,
    }))
